#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""doc2md.py — 原件（DOCX / 有文字层 PDF / PPTX）→ 题级 Markdown 候选的确定性转换。

用途
  把一份原件按确定性规则转成与题库现有题级 MD 同一套小节标题的候选稿（bank.py 能直接读），
  每页/每块标注页型（text_reliable / mixed / scanned / garbled_font / empty）与来源页码；
  题号、小问、分值、选项、表格、图片占位按原文结构还原；脚本做不了的部分（图、扫描页、乱码字体页、
  OCR 文字层页〔含 Adobe ClearScan〕、浮动文本框顺序、复杂表格、切题存疑、被剔除的不可见文字……）
  单列 residual 清单交模型看图。
  扫描页只做页型判定和 OCR 候选标注（已有 OCR 或 --run-ocr），一律不当定稿。
  doc/rtf 旧格式先用 macOS 自带 textutil 转 docx 再走 DOCX 路径（标注 legacy_converted）。

用法（必须用 /usr/bin/python3〔3.9〕，依赖 PyMuPDF 与 python-pptx）
  单个原件：
    python3 doc2md.py 原件.pdf --role 原卷 [--exam-id BJ-2025-DC-ERMO] --out-dir DIR
  一卷全部来源（按题库转换主清单 conversion_manifest.json 的 tasks[].sources，合并出题级候选）：
    python3 doc2md.py --exam BJ-2025-DC-ERMO --out-dir DIR [--run-ocr]
  作为库：
    from doc2md import extract_source, split_paper, split_rubric, convert_file, convert_exam, TextIndex
输出（全部在 --out-dir 下；不给 --out-dir 时取环境变量 DOC2MD_OUT，两者都没有则拒绝运行）
  单件：<名>/source.full.md 逐页全文、<名>/questions/<QID>.md 题级候选、pages.jsonl 页型、
        residual.jsonl 残留、selfcheck.json 自检、assets/ 嵌图与裁图
  整卷：<EXAM>/<QID>.md（合并题面/E1/E3/讲评各层）、<EXAM>/sources/<sid>/…、
        <EXAM>/selfcheck.json、<EXAM>/residual.jsonl、<EXAM>/pages.jsonl、00_exam_front_matter.md
        题面来源在原卷/教师版之间按“切题完整度 → 文字层占比 → 角色”排序选取（selfcheck.paper_candidates 可查）。
        selfcheck.routes 给每题分流建议：A 纯文字层无残留；B 有图/表/框/文本框须看；
        C 题面只有 OCR/扫描/乱码、切题存疑、或本题/整卷有硬失败，须模型对照原页。

自检硬失败（任一条即退出码 1）
  题号不连续；题数与主清单 question_index_count / question_md_count / 题库题文件数都不符；
  题数与卷首“本部分共N题”合计不符（分部总分合计为 100 时）；选择题选项不齐或字母重复；
  分值有未解析题且卷首分部总分不是 100；逐题分值合计≠100；答案块串题（一块里有别题的
  “N．【分析】”题首或多个“故选”）；本件有同类题答案块而某题缺；图片引用失效；原件缺失或打不开。

只读承诺
  原件、题库（DeepSeek_政治题库资料库_20260918）、书稿、审阅入口、协作文件一律只读；输出目录落在
  题库、原料根、任一“YYYY模拟题”原料目录、原件所在目录、桌面（Luna 写权区除外，须
  BANK_ACTOR=shared_writer）、Skill 目录或任一书册 profile 的 paths.* 目录内时直接拒绝；
  落在冻结册（profile.frozen=true）范围时调用 frozen_guard。路径比较按 inode 与大小写折叠双重判断。
  脚本只截取原文，不生成归纳或摘要；截断处一律标“…〔截断…〕”。
  所有产物证据等级为 machine_extracted，不得标为已核验。
  机器自检全过只说明机械层没有发现问题，不等于内容已核验。

退出码
  0 转换完成且自检无硬失败；1 转换完成、已写 selfcheck.json 且有硬失败；
  2 参数/输入错误、原件打不开、依赖缺失、拒绝写入或其他异常（不写或只写部分产物）。
"""
import argparse
import collections
import csv
import hashlib
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile
import xml.etree.ElementTree as ET

sys.dont_write_bytecode = True  # 不在 Skill 目录与题库 scripts 目录留 .pyc
VERSION = "doc2md-2026.09.24a"
HERE = os.path.dirname(os.path.abspath(__file__))
CANON_ROOT = "/Users/wanglifei/Desktop/gpt和claude共同的小窝"
# BAODIAN_ROOT 只改“读哪里”，不缩小护栏：护栏同时保护 CANON_ROOT 与 BAODIAN_ROOT 两套路径
ROOT = os.environ.get("BAODIAN_ROOT", CANON_ROOT)
BANK = os.path.join(ROOT, "DeepSeek_政治题库资料库_20260918")
RAW = os.path.join(ROOT, "00_共同资料/原材料")
DESKTOP = os.path.dirname(ROOT)
PIPE = os.path.join(ROOT, "后勤管理/MD全库流水线_20260921")
MANIFEST = os.path.join(PIPE, "conversion_manifest.json")
STATE = os.path.join(PIPE, "controller_transition_20260922/state.json")
# 输出根：只认 --out-dir 或环境变量 DOC2MD_OUT；不写死任何会话临时目录
DEFAULT_OUT = os.environ.get("DOC2MD_OUT") or None
RE_SOURCE_TOP = re.compile(r"^(\d{4}模拟题|历年高考题及细则|宝典制作原料)$")

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
WPS = "{http://schemas.microsoft.com/office/word/2010/wordprocessingShape}"
V = "{urn:schemas-microsoft-com:vml}"
PR = "{http://schemas.openxmlformats.org/package/2006/relationships}"

CN_NUM = "〇一二三四五六七八九十"
CIRCLED = "①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳"
TERMINAL = "。！？；：…”」』）)】"
DOTS = "．.·•・"

# 细页型 → 任务要求的粗页型
COARSE = {
    "text_reliable": "text_reliable", "docx": "text_reliable", "pptx_text": "text_reliable",
    "text_plus_figure": "mixed", "docx_with_image": "mixed", "pptx_mixed": "mixed",
    "legacy_converted": "mixed",
    "ocr_layer_over_scan": "scanned", "scan_with_ocr_layer": "scanned", "scan_no_text": "scanned",
    "ocr_text_layer": "scanned",
    "vector_outlined_text": "scanned", "pptx_picture": "scanned",
    "symbol_garbled": "garbled_font", "garbled_text_layer": "garbled_font",
    "empty": "empty",
}
# 可当作 native 文字直出的细页型；其余只能是 OCR 候选或需看图
DIRECT_CLASSES = {"text_reliable", "text_plus_figure", "docx", "docx_with_image", "pptx_text", "pptx_mixed",
                  "legacy_converted"}
OCR_LAYER_CLASSES = {"ocr_layer_over_scan", "scan_with_ocr_layer", "ocr_text_layer"}

ROLE_LAYER = {"原卷": "paper", "教师版": "teacher", "混合载体": "teacher", "教师版/转录题源": "teacher",
              "评分材料": "E1", "评分细则_分题切片": "E1", "阅卷总结": "E1",
              "参考答案": "E3", "讲评": "lecture", "试题分析材料": "lecture"}
LAYER_TITLE = {
    "paper": "题目原文（原卷·native，读取优先）",
    "teacher": "题面（E0｜教师版题面来源）",
    "ocr": "原卷 OCR 候选（未逐字对照，不得直接采用）",
    "E1": "正式评分材料原文",
    "E3": "参考答案原文",
    "lecture": "讲评来源块（含答案与解析，不得当作题面）",
    "aux": "待确认材料来源块（角色未定，须人工裁决）",
}


# ---------------------------------------------------------------- 通用
def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for b in iter(lambda: fh.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def sid_of(path, digest=None):
    """题库 source_id 规则：S + sha256 前 12 位。"""
    return "S" + (digest or sha256(path))[:12]


def cn(n):
    if n <= 10:
        return CN_NUM[n] if n else "〇"
    if n < 20:
        return "十" + (CN_NUM[n - 10] if n > 10 else "")
    return CN_NUM[n // 10] + "十" + (CN_NUM[n % 10] if n % 10 else "")


def roman(n):
    vals = [(1000, "M"), (900, "CM"), (500, "D"), (400, "CD"), (100, "C"), (90, "XC"), (50, "L"), (40, "XL"),
            (10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")]
    s = ""
    for v, r in vals:
        while n >= v:
            s += r
            n -= v
    return s


def fmt_num(n, fmt):
    if fmt in ("decimal", "decimalZero"):
        return str(n)
    if fmt == "upperLetter":
        return chr(64 + n) if 1 <= n <= 26 else str(n)
    if fmt == "lowerLetter":
        return chr(96 + n) if 1 <= n <= 26 else str(n)
    if fmt == "upperRoman":
        return roman(n)
    if fmt == "lowerRoman":
        return roman(n).lower()
    if fmt in ("decimalEnclosedCircle", "decimalEnclosedCircleChinese"):
        return CIRCLED[n - 1] if 1 <= n <= 20 else str(n)
    if fmt == "decimalFullWidth":
        return "".join(chr(0xFF10 + int(c)) for c in str(n))
    if fmt in ("chineseCounting", "chineseCountingThousand", "ideographTraditional", "japaneseCounting",
               "ideographDigital", "chineseLegalSimplified"):
        return cn(n)
    return str(n)


def plain(t):
    """去掉本脚本加的内联标记，供正则判断用（不改输出文字）。"""
    t = re.sub(r"<[^>]+>|\*\*", "", t or "")
    t = re.sub(r"〔(OCR候选|文本框|字体映射疑误)〕", "", t)
    return t.lstrip("> ")


def clip(t, n, ref=None):
    """截取前 n 字；超长时在截断处标注（红线：脚本只截取原文，截断处必须标出）。"""
    t = t or ""
    if len(t) <= n:
        return t
    return t[:n] + ("…〔截断，全文见 %s〕" % ref if ref else "…〔截断〕")


def clip_window(t, a, b, ref=None):
    """取 t[a:b] 作上下文窗口，左右被截处分别标注。返回 (文字, 是否截断)。"""
    a, b = max(0, a), min(len(t), b)
    s = t[a:b]
    cut = False
    if a > 0:
        s = "〔前略〕" + s
        cut = True
    if b < len(t):
        s = s + ("〔后略，全文见 %s〕" % ref if ref else "〔后略〕")
        cut = True
    return s, cut


# ---------------------------------------------------------------- DOCX
class Docx:
    """按 body 顺序解析 document.xml：自动编号、文本框、嵌图、表格、公式、修订、着重号。"""

    def __init__(self, path, sid, asset_root=None):
        self.z = zipfile.ZipFile(path)
        self.sid = sid
        self.asset_root = asset_root
        self.rels = self._rels("word/_rels/document.xml.rels")
        self.num = self._numbering()
        self.style_num = self._style_numbering()
        self.counters = collections.defaultdict(lambda: [0] * 9)
        self.blocks = []
        self.page = 1
        self.img_i = 0
        self.flags = collections.Counter()
        self._has_text = False

    def _xml(self, name):
        try:
            return ET.fromstring(self.z.read(name))
        except KeyError:
            return None

    def _rels(self, name):
        x = self._xml(name)
        out = {}
        if x is not None:
            for r in x.findall(PR + "Relationship"):
                out[r.get("Id")] = r.get("Target")
        return out

    def _numbering(self):
        x = self._xml("word/numbering.xml")
        absn, nums = {}, {}
        if x is None:
            return {}
        for an in x.findall(W + "abstractNum"):
            lv = {}
            for l in an.findall(W + "lvl"):
                il = int(l.get(W + "ilvl"))

                def g(t, d=None, l=l):
                    e = l.find(W + t)
                    return e.get(W + "val") if e is not None else d
                lv[il] = {"fmt": g("numFmt", "decimal"), "text": g("lvlText", ""), "start": int(g("start", "1"))}
            absn[an.get(W + "abstractNumId")] = lv
        for n in x.findall(W + "num"):
            ai = n.find(W + "abstractNumId")
            if ai is None:
                continue
            lv = {k: dict(v) for k, v in absn.get(ai.get(W + "val"), {}).items()}
            for ov in n.findall(W + "lvlOverride"):
                il = int(ov.get(W + "ilvl"))
                so = ov.find(W + "startOverride")
                if so is not None and il in lv:
                    lv[il]["start"] = int(so.get(W + "val"))
            nums[n.get(W + "numId")] = lv
        return nums

    def _style_numbering(self):
        x = self._xml("word/styles.xml")
        out = {}
        if x is None:
            return out
        for st in x.findall(W + "style"):
            np_ = st.find(W + "pPr/" + W + "numPr")
            if np_ is not None:
                nid = np_.find(W + "numId")
                il = np_.find(W + "ilvl")
                out[st.get(W + "styleId")] = (nid.get(W + "val") if nid is not None else None,
                                              int(il.get(W + "val")) if il is not None else 0)
        return out

    def num_label(self, p):
        ppr = p.find(W + "pPr")
        nid, il = None, 0
        if ppr is not None:
            np_ = ppr.find(W + "numPr")
            ps = ppr.find(W + "pStyle")
            if ps is not None and ps.get(W + "val") in self.style_num:
                nid, il = self.style_num[ps.get(W + "val")]
            if np_ is not None:
                a = np_.find(W + "numId")
                b = np_.find(W + "ilvl")
                if a is not None:
                    nid = a.get(W + "val")
                if b is not None:
                    il = int(b.get(W + "val"))
        if not nid or nid == "0" or nid not in self.num:
            return ""
        lv = self.num[nid]
        if il not in lv:
            return ""
        c = self.counters[nid]
        if c[il] == 0:
            c[il] = lv[il]["start"]
        else:
            c[il] += 1
        for d in range(il + 1, 9):
            c[d] = 0
        if lv[il]["fmt"] == "bullet":
            self.flags["bullet"] += 1
            return "• "
        t = lv[il]["text"]
        for k in range(1, 10):
            if "%%%d" % k in t:
                lvk = lv.get(k - 1, {"fmt": "decimal", "start": 1})
                val = c[k - 1] or lvk["start"]
                t = t.replace("%%%d" % k, fmt_num(val, lvk["fmt"]))
        self.flags["auto_number"] += 1
        return t

    def save_media(self, rid):
        tgt = self.rels.get(rid)
        if not tgt:
            return None
        name = tgt[1:] if tgt.startswith("/") else "word/" + tgt
        try:
            data = self.z.read(os.path.normpath(name))
        except KeyError:
            return None
        self.img_i += 1
        ext = os.path.splitext(tgt)[1] or ".bin"
        rel = "assets/%s/img%02d%s" % (self.sid, self.img_i, ext)
        if self.asset_root:
            full = os.path.join(self.asset_root, rel)
            os.makedirs(os.path.dirname(full), exist_ok=True)
            with open(full, "wb") as fh:
                fh.write(data)
        return rel

    def inline(self, el, out, extra):
        for ch in el:
            tag = ch.tag
            if tag == MC + "AlternateContent":
                choice = ch.find(MC + "Choice")  # 只取 Choice，Fallback 是同一内容的旧格式副本
                if choice is not None:
                    self.inline(choice, out, extra)
                continue
            if tag == W + "r":
                self.run(ch, out, extra)
            elif tag in (W + "hyperlink", W + "smartTag", W + "ins", W + "fldSimple", W + "sdt", W + "sdtContent",
                         W + "customXml"):
                if tag == W + "ins":
                    self.flags["tracked_insert"] += 1
                self.inline(ch, out, extra)
            elif tag == W + "del":
                txt = "".join(t.text or "" for t in ch.iter(W + "delText"))
                if txt:
                    out.append("〔修订删除:%s〕" % txt)
                    self.flags["tracked_delete"] += 1
            elif tag in (M + "oMath", M + "oMathPara"):
                out.append("$" + "".join(t.text or "" for t in ch.iter(M + "t")) + "$")
                self.flags["formula"] += 1

    def run(self, r, out, extra):
        rpr = r.find(W + "rPr")
        b = u = sup = sub = em = hidden = False
        mark, color = False, None
        if rpr is not None:
            bb = rpr.find(W + "b")
            b = bb is not None and bb.get(W + "val") not in ("0", "false")
            uu = rpr.find(W + "u")
            u = uu is not None and uu.get(W + "val") != "none"
            va = rpr.find(W + "vertAlign")
            if va is not None:
                sup = va.get(W + "val") == "superscript"
                sub = va.get(W + "val") == "subscript"
            ee = rpr.find(W + "em")
            em = ee is not None and ee.get(W + "val") not in ("none", None)
            hl = rpr.find(W + "highlight")
            shd = rpr.find(W + "shd")
            mark = (hl is not None and hl.get(W + "val") not in ("none", None)) or \
                   (shd is not None and shd.get(W + "fill") not in (None, "auto", "FFFFFF"))
            col = rpr.find(W + "color")
            color = col.get(W + "val") if col is not None and col.get(W + "val") not in (None, "auto", "000000") else None
            vn = rpr.find(W + "vanish")
            hidden = vn is not None and vn.get(W + "val") not in ("0", "false")
        buf = []
        for ch in r:
            tag = ch.tag
            if tag == W + "t":
                buf.append(ch.text or "")
                if (ch.text or "").strip():
                    self._has_text = True
            elif tag == W + "tab":
                buf.append("\t")
            elif tag in (W + "br", W + "cr"):
                if ch.get(W + "type") == "page":
                    self.page += 1
                else:
                    buf.append("\n")
            elif tag == W + "lastRenderedPageBreak":
                self.page += 1
            elif tag == W + "sym":
                code = (ch.get(W + "char", "") or "").upper()
                mp = {"F0FC": "✓", "F0FB": "✗", "F0A3": "□", "F0A8": "□", "F06F": "□", "F0B7": "•", "F0D8": "▲",
                      "F0E0": "→", "F0DF": "←", "F081": "①", "F082": "②", "F083": "③", "F084": "④"}
                buf.append(mp.get(code, "〔符号:%s:%s〕" % (ch.get(W + "font", ""), code)))
                self.flags["sym"] += 1
            elif tag == W + "noBreakHyphen":
                buf.append("-")
            elif tag == MC + "AlternateContent":
                choice = ch.find(MC + "Choice")
                if choice is not None:
                    self.drawing(choice, extra)
            elif tag in (W + "drawing", W + "pict", W + "object"):
                self.drawing(ch, extra)
        s = "".join(buf)
        if not s:
            return
        if hidden:
            self.flags["hidden_text"] += 1
            out.append("〔隐藏文字:%s〕" % s)
            return
        if sup:
            s = "<sup>%s</sup>" % s
        elif sub:
            s = "<sub>%s</sub>" % s
        if u and s.strip():
            s = "<u>%s</u>" % s
        if mark and s.strip():
            s = "<mark>%s</mark>" % s
            self.flags["highlight"] += 1
        if color and s.strip():
            s = '<span data-color="%s">%s</span>' % (color, s)
            self.flags["color"] += 1
        if em and s.strip():
            s = '<em class="dot">%s</em>' % s
            self.flags["emphasis_dot"] += 1
        if b and s.strip():
            s = "**%s**" % s
        out.append(s)

    def drawing(self, el, extra):
        tb = list(el.iter(W + "txbxContent"))
        for t in tb:
            paras = []
            for p in t.iter(W + "p"):
                o = []
                self.inline(p, o, [])
                s = (self.num_label(p) + "".join(o)).strip()
                if s:
                    paras.append(s)
            if paras:
                extra.append({"kind": "textbox", "text": "\n".join(paras), "page": self.page,
                              "before": not self._has_text})
                self.flags["textbox"] += 1
        blips = list(el.iter(A + "blip")) + list(el.iter(V + "imagedata"))
        for bl in blips:
            rid = bl.get(R + "embed") or bl.get(R + "id")
            rel = self.save_media(rid)
            extra.append({"kind": "image", "text": "", "path": rel, "page": self.page})
            self.flags["image"] += 1
        if not tb and not blips and (list(el.iter(WPS + "wsp")) or list(el.iter(V + "shape")) or
                                     list(el.iter(V + "line")) or list(el.iter(V + "rect"))):
            extra.append({"kind": "shape", "text": "", "page": self.page})
            self.flags["shape_no_text"] += 1

    def para(self, p):
        out, extra = [], []
        lab = self.num_label(p)
        page0 = self.page
        saved = self._has_text
        self._has_text = False
        self.inline(p, out, extra)
        self._has_text = saved
        s = lab + "".join(out)
        for _ in range(3):
            s = re.sub(r"</(u|mark|sup|sub)><\1>", "", s)
            s = re.sub(r'</em><em class="dot">', "", s)
            s = s.replace("****", "")
        res = [b for b in extra if b.get("before")]
        if s.strip():
            res.append({"kind": "p", "text": s.rstrip(), "page": page0})
        res.extend(b for b in extra if not b.get("before"))
        return res

    def table(self, tbl):
        rows, merged = [], False
        for tr in tbl.findall(W + "tr"):
            cells = []
            for tc in tr.findall(W + "tc"):
                tcpr = tc.find(W + "tcPr")
                span, vm = 1, None
                if tcpr is not None:
                    gs = tcpr.find(W + "gridSpan")
                    if gs is not None:
                        span = int(gs.get(W + "val"))
                        merged = merged or span > 1
                    v = tcpr.find(W + "vMerge")
                    if v is not None:
                        vm = v.get(W + "val") or "continue"
                        merged = True
                parts = []
                for ch in tc:
                    if ch.tag == W + "p":
                        for b in self.para(ch):
                            if b["kind"] in ("p", "textbox"):
                                parts.append(b["text"])
                            elif b["kind"] == "image":
                                parts.append("![图](%s)" % b["path"])
                                self.blocks.append(dict(b, in_table=True))
                    elif ch.tag == W + "tbl":
                        parts.append("〔嵌套表格〕" + self.table(ch)["text"].replace("\n", " "))
                cells.append({"text": "<br>".join(x.replace("\n", "<br>") for x in parts), "span": span, "vm": vm})
            rows.append(cells)
        self.flags["table"] += 1
        return {"kind": "table", "text": table_md(rows, merged), "page": self.page, "rows": len(rows),
                "cols": table_shape(rows), "merged": merged}

    def extract(self):
        body = self._xml("word/document.xml").find(W + "body")
        for ch in body:
            if ch.tag == W + "p":
                self.blocks.extend(self.para(ch))
            elif ch.tag == W + "tbl":
                self.blocks.append(self.table(ch))
            elif ch.tag in (W + "sdt", MC + "AlternateContent"):
                for p in ch.iter(W + "p"):
                    self.blocks.extend(self.para(p))
        return self.blocks


def table_shape(rows):
    """每行的逻辑列数（按 gridSpan 展开）。"""
    return [sum(c.get("span", 1) for c in r) for r in rows]


def table_md(rows, merged):
    if not rows:
        return ""
    if merged:
        h = ["<table>"]
        for r in rows:
            h.append("<tr>" + "".join(
                "<td%s%s>%s</td>" % (' colspan="%d"' % c["span"] if c["span"] > 1 else "",
                                     ' data-vmerge="%s"' % c["vm"] if c["vm"] else "",
                                     "" if c["vm"] == "continue" else c["text"]) for c in r) + "</tr>")
        h.append("</table>")
        return "\n".join(h)
    ncol = max(len(r) for r in rows)
    lines = []
    for i, r in enumerate(rows):
        cells = [c["text"].replace("|", "\\|").replace("\n", "<br>") for c in r] + [""] * (ncol - len(r))
        lines.append("| " + " | ".join(cells) + " |")
        if i == 0:
            lines.append("|" + "---|" * ncol)
    return "\n".join(lines)


def docx_page_class(blocks):
    """DOCX 没有物理页型问题：按估算页标 docx / docx_with_image。"""
    pages = collections.OrderedDict()
    for b in blocks:
        pg = b.get("page", 1)
        d = pages.setdefault(pg, {"page": pg, "chars": 0, "images": 0, "textboxes": 0})
        if b["kind"] == "image":
            d["images"] += 1
        elif b["kind"] == "textbox":
            d["textboxes"] += 1
        d["chars"] += len(re.sub(r"\s", "", plain(b.get("text", ""))))
    out = []
    for pg, d in pages.items():
        cls = "docx_with_image" if d["images"] else "docx"
        out.append(dict(d, fine_class=cls, page_type=COARSE[cls], page_basis="lastRenderedPageBreak估算"))
    return out


# ---------------------------------------------------------------- PPTX
def open_pptx(path):
    from pptx import Presentation
    try:
        return Presentation(path), None
    except Exception as e:
        # D-002/D-003 类：zip 结构不被接受时在内存里重打包再试（不改原件）
        try:
            src = zipfile.ZipFile(path)
            buf = io.BytesIO()
            with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
                for info in src.infolist():
                    dst.writestr(info.filename, src.read(info.filename))
            buf.seek(0)
            return Presentation(buf), "repacked_in_memory:%s" % type(e).__name__
        except Exception as e2:
            raise RuntimeError("PPTX 打不开：%r / %r" % (e, e2))


PPTX_AUTONUM = {"arabicPeriod": "{n}.", "arabicParenR": "{n})", "arabicParenBoth": "({n})", "arabicPlain": "{n}",
                "circleNumDbPlain": "{c}", "circleNumWdBlackPlain": "{c}", "circleNumWdWhitePlain": "{c}",
                "romanUcPeriod": "{R}.", "romanLcPeriod": "{r}.", "alphaUcPeriod": "{A}.", "alphaLcPeriod": "{a}.",
                "alphaLcParenR": "{a})", "alphaUcParenR": "{A})", "ea1ChsPeriod": "{z}.", "ea1ChsPlain": "{z}、",
                "ea1JpnChsDbPeriod": "{z}."}


def _pptx_autonum(p, counters):
    """段落的 a:buAutoNum 自动编号 → 文字标签（与 PowerPoint 显示一致的常见格式）；不是自动编号返回空串。"""
    ppr = p._p.pPr
    lvl = p.level or 0
    if ppr is None:
        counters.pop(lvl, None)
        return ""
    au = ppr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}buAutoNum")
    if au is None or not (p.text or "").strip():
        if ppr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}buNone") is not None:
            counters.pop(lvl, None)
        return ""
    typ = au.get("type", "arabicPeriod")
    start = int(au.get("startAt", "1"))
    key = (lvl, typ)
    n = counters.get(key, start - 1) + 1
    counters[key] = n
    fmt = PPTX_AUTONUM.get(typ, "{n}.")
    return fmt.format(n=n, c=CIRCLED[n - 1] if 1 <= n <= 20 else n, R=roman(n), r=roman(n).lower(),
                      A=chr(64 + n) if n <= 26 else n, a=chr(96 + n) if n <= 26 else n, z=cn(n))


def extract_pptx(path, sid, asset_root=None):
    prs, note = open_pptx(path)
    blocks, pages = [], []
    flags = collections.Counter()
    if note:
        flags[note] += 1
    img_i = [0]
    slide_area = float((prs.slide_width or 1) * (prs.slide_height or 1))

    def walk(shapes, page, stat):
        items = []
        for sh in shapes:
            try:
                top, left = sh.top or 0, sh.left or 0
            except Exception:
                top = left = 0
            items.append((top, left, sh))
        for top, left, sh in sorted(items, key=lambda x: (x[0] // 200000, x[1])):
            st = getattr(sh, "shape_type", None)
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                # 一段一块：文本框里“参考答案：/17.（8分）/…”这样的题号行才能被切题识别；自动编号（buAutoNum）按级计数补出
                counters = {}
                for p in sh.text_frame.paragraphs:
                    t = "".join(r.text for r in p.runs) or p.text
                    lab = _pptx_autonum(p, counters)
                    if t.strip():
                        blocks.append({"kind": "p", "text": ("  " * (p.level or 0)) + lab + t, "page": page})
                        stat["chars"] += len(re.sub(r"\s", "", t))
                        if lab:
                            flags["pptx_auto_number"] += 1
            if getattr(sh, "has_table", False) and sh.has_table:
                rows = [[{"text": c.text.replace("\n", "<br>"), "span": 1, "vm": None} for c in r.cells]
                        for r in sh.table.rows]
                blocks.append({"kind": "table", "text": table_md(rows, False), "page": page, "rows": len(rows),
                               "cols": table_shape(rows), "merged": False})
                stat["chars"] += sum(len(re.sub(r"\s", "", c["text"])) for r in rows for c in r)
                flags["table"] += 1
            if getattr(sh, "has_chart", False) and sh.has_chart:
                blocks.append({"kind": "shape", "text": "", "page": page, "what": "chart"})
                stat["figures"] += 1
                flags["chart"] += 1
            el = getattr(sh, "_element", None)
            if el is not None and el.tag.endswith("graphicFrame") and b"dgm" in ET.tostring(el)[:4000]:
                blocks.append({"kind": "shape", "text": "", "page": page, "what": "smartart"})
                flags["smartart"] += 1
            if st == 13 or sh.__class__.__name__ == "Picture":
                try:
                    img = sh.image
                    img_i[0] += 1
                    rel = "assets/%s/s%03d_img%02d.%s" % (sid, page, img_i[0], img.ext)
                    if asset_root:
                        full = os.path.join(asset_root, rel)
                        os.makedirs(os.path.dirname(full), exist_ok=True)
                        with open(full, "wb") as fh:
                            fh.write(img.blob)
                    area = (sh.width or 0) * (sh.height or 0) / slide_area
                    blocks.append({"kind": "image", "text": "", "path": rel, "page": page, "area": round(area, 3)})
                    stat["pic_area"] += area
                    stat["figures"] += 1
                    flags["image"] += 1
                except Exception:
                    flags["image_unreadable"] += 1
            if st == 6 or sh.__class__.__name__ == "GroupShape":
                walk(sh.shapes, page, stat)

    for i, s in enumerate(prs.slides, 1):
        hidden = s._element.get("show") == "0"
        blocks.append({"kind": "pagemark", "text": "", "page": i, "hidden": hidden})
        stat = collections.Counter()
        walk(s.shapes, i, stat)
        if s.has_notes_slide and s.notes_slide.notes_text_frame is not None:
            nt = s.notes_slide.notes_text_frame.text
            if nt.strip():
                blocks.append({"kind": "notes", "text": nt, "page": i})
        if stat["pic_area"] >= 0.5 and stat["chars"] < 30:
            cls = "pptx_picture"
        elif stat["figures"]:
            cls = "pptx_mixed"
        elif stat["chars"] == 0:
            cls = "empty"
        else:
            cls = "pptx_text"
        for b in blocks:
            if b.get("page") == i and b["kind"] == "pagemark":
                b["class"] = cls
        pages.append({"page": i, "fine_class": cls, "page_type": COARSE[cls], "chars": stat["chars"],
                      "pic_area": round(stat["pic_area"], 3), "hidden": hidden})
    return blocks, flags, pages


# ---------------------------------------------------------------- PDF
RE_PAGENO = re.compile(r"^(第\s*\d+\s*页|[-—]?\s*\d{1,3}\s*[-—]?|.{0,30}第\s*\d+\s*页\s*[（(]?\s*共\s*\d+\s*页[）)]?|"
                       r"共\s*\d+\s*页.*|\d+\s*/\s*\d+)$")
RE_QSTART = re.compile(r"^\s*(\d{1,2})\s*[．.、](?!\d+(?:\.\d+)?\s*[%％万亿倍个])\s*")
RE_OPT = re.compile(r"^\s*[A-DＡ-Ｄ]\s*[．.、]")
RE_SUBQ = re.compile(r"^\s*[（(]\s*\d{1,2}\s*[）)]")
RE_CIRC = re.compile(r"^\s*[①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳]")
RE_MAT = re.compile(r"^\s*(材料[一二三四五六七八九十]|【|注[:：]|第[一二三]部分|示例|细则|评分|说明|本部分)")
CJKP = "一-鿿，。、；：？！“”‘’（）《》【】—…"
SYM = re.compile(r'[!"#$%&*+]')
CJK = re.compile(r"[一-鿿㐀-䶿]")
# 乱码判据：替换符、控制符、私用区（F000–F0FF 是 Symbol/Wingdings 符号字体的正常映射，不算）、彝文音节、补充私用区
BAD = re.compile("[\ufffd\x00-\x08\x0b\x0c\x0e-\x1f\ue000-\uefff\uf100-\uf8ff\ua000-\ua4cf\U000f0000-\U0010ffff]")
# 康熙部首/兼容汉字：个别 PDF 生成器把“用人一力”写成 U+2F00 段的部首字符，字形相同、码位不同
RADICALS = re.compile("[\u2e80-\u2fdf\uf900-\ufaff]")


RADICAL_SUPP = dict(zip("⺁⺇⺌⺗⺘⺝⺟⺠⺢⺣⺧⺩⺫⺬⺮⺱⺶⺼⺽⻁⻄⻅⻆⻉⻊⻋⻑⻓⻔⻗⻘⻙⻚⻛⻜⻝⻢⻣⻤⻥⻦⻧⻨⻩⻫⻬⻭⻯⻰⻱⻳",
                        "厂几小心手月母民水火牛王目示竹网羊肉臼虎西见角贝足车長长门雨青韦页风飞食马骨鬼鱼鸟卤麦黄齐齐齿龙龙龟龟"))


def _radical(ch):
    import unicodedata
    if ch in RADICAL_SUPP:
        return RADICAL_SUPP[ch]
    n = unicodedata.normalize("NFKC", ch)
    return n if len(n) == 1 else ch


# 方正类排版软件导出的 PDF 把中文标点映射到彝文音节/补充私用区（2024朝阳一模、2026顺义一模实测，对照题库人工稿确认）
FOUNDER_PUNCT = {"\ua3ac": "，", "\ua3ae": "。", "\ua3bb": "；", "\U001001b0": "．", "\U001001ba": "—"}
RE_FOUNDER = re.compile("|".join(re.escape(k) for k in FOUNDER_PUNCT))


def fix_founder_punct(t):
    return RE_FOUNDER.sub(lambda m: FOUNDER_PUNCT[m.group(0)], t)


def fix_radicals(t):
    """把康熙部首/部首补充/兼容汉字码位换成统一汉字（一字对一字）；只换码位不换字形。"""
    return RADICALS.sub(lambda m: _radical(m.group(0)), t)


def squeeze(t):
    """PDF 文字层在中文旁插入的排版空格：相邻任一侧为中文或全角标点即删除。"""
    return re.sub(r"(?<=[%s]) +| +(?=[%s])" % (CJKP, CJKP), "", t)


def _coverage(rects, W_, H_, G=40):
    if not rects:
        return 0.0
    grid = [[0] * G for _ in range(G)]
    for r in rects:
        x0 = max(0, min(G, int(r.x0 / W_ * G)))
        x1 = max(0, min(G, -int(-r.x1 / W_ * G // 1)))
        y0 = max(0, min(G, int(r.y0 / H_ * G)))
        y1 = max(0, min(G, -int(-r.y1 / H_ * G // 1)))
        for y in range(y0, y1):
            for x in range(x0, x1):
                grid[y][x] = 1
    return sum(map(sum, grid)) / float(G * G)


# 印厂排版标签（出血区的 .indd 文件名与输出时间），不是正文
RE_TYPESET_LABEL = re.compile(r"\.indd\b|^\s*\d{4}/\d{1,2}/\d{1,2}\s+\d{1,2}:\d{2}(?::\d{2})?\s*$", re.I)
# Adobe ClearScan / Paper Capture 合成字体（“Adobe Song Std L-7528”“SimSun-7526-Identity-H”）：印的是扫描字形，文字层是 OCR
RE_CLEARSCAN_FONT = re.compile(r"Std L-(?:Bold-)?\d+|-\d{3,5}-Identity-[HV]$")
# OCR 把 ①②③④ 认成 CD/©/＠/® 之类：行首替代符后紧跟汉字
RE_OCR_CIRCLE_SUBST = re.compile(r"^\s*(?:CD|©|＠|®|@|◎|\(D)\s*[一-鿿]")


def _is_white(color):
    r, g, b_ = (color >> 16) & 255, (color >> 8) & 255, color & 255
    return min(r, g, b_) >= 0xF0


class _PageRaster:
    """按需低清渲染一页，用来判断白字底下是不是白底、页面墨迹是否在文字层之外。"""

    def __init__(self, pg, dpi=36):
        self.pg, self.dpi, self._pix = pg, dpi, None

    def pix(self):
        if self._pix is None:
            import fitz
            self._pix = self.pg.get_pixmap(dpi=self.dpi, colorspace=fitz.csGRAY, alpha=False)
        return self._pix

    def _box(self, r):
        px = self.pix()
        k = self.dpi / 72.0
        x0 = max(0, min(px.width - 1, int(r[0] * k)))
        y0 = max(0, min(px.height - 1, int(r[1] * k)))
        x1 = max(x0 + 1, min(px.width, int(r[2] * k + 0.999)))
        y1 = max(y0 + 1, min(px.height, int(r[3] * k + 0.999)))
        return px, x0, y0, x1, y1

    def nonwhite_ratio(self, r, thr=225):
        px, x0, y0, x1, y1 = self._box(r)
        buf, st = px.samples, px.stride
        n = hit = 0
        for y in range(y0, y1):
            row = buf[y * st + x0:y * st + x1]
            n += len(row)
            hit += sum(1 for v in row if v < thr)
        return hit / float(n or 1)

    def ink_outside(self, rects, thr=200):
        """页面墨迹（灰度<thr）落在给定文字框之外的像素占全页比例。"""
        px = self.pix()
        k = self.dpi / 72.0
        W_, H_ = px.width, px.height
        mask = bytearray(W_ * H_)
        for r in rects:
            x0, y0 = max(0, int(r[0] * k) - 1), max(0, int(r[1] * k) - 1)
            x1, y1 = min(W_, int(r[2] * k) + 2), min(H_, int(r[3] * k) + 2)
            for y in range(y0, y1):
                mask[y * W_ + x0:y * W_ + x1] = b"\x01" * (x1 - x0)
        buf, st = px.samples, px.stride
        hit = 0
        for y in range(H_):
            row = buf[y * st:y * st + W_]
            m = mask[y * W_:(y + 1) * W_]
            hit += sum(1 for v, mm in zip(row, m) if v < thr and not mm)
        return hit / float(W_ * H_ or 1)


def page_spans(pg, raster=None):
    """rawdict 各行的 span，逐个判可见性。返回 (lines[(line, kept_spans)], dropped[], stat)。
    剔除：alpha<128 的水印；白字且其下为白底（渲染像素判断）；印厂排版标签行。
    深色填充或图片上的白字是可见字，保留（stat['white_on_dark_kept'] 计数）。凡剔除的都进 dropped，供 residual 列出。"""
    raster = raster or _PageRaster(pg)
    out, dropped = [], []
    stat = collections.Counter()
    for b in pg.get_text("rawdict", sort=True).get("blocks", []):
        if b.get("type") != 0:
            continue
        for l in b.get("lines", []):
            kept = []
            for sp in l.get("spans", []):
                if not sp.get("chars"):
                    continue
                txt = "".join(c["c"] for c in sp["chars"])
                if not txt.strip():
                    kept.append(sp)
                    continue
                if sp.get("alpha", 255) < 128:
                    dropped.append({"reason": "low_alpha_watermark", "text": txt, "bbox": [round(v, 1) for v in sp["bbox"]]})
                    stat["dropped_low_alpha"] += 1
                    continue
                if _is_white(sp.get("color", 0)):
                    ratio = raster.nonwhite_ratio(sp["bbox"])
                    if ratio < 0.35:
                        dropped.append({"reason": "white_on_white", "text": txt, "bbox": [round(v, 1) for v in sp["bbox"]],
                                        "bg_nonwhite_ratio": round(ratio, 3)})
                        stat["dropped_white_on_white"] += 1
                        continue
                    stat["white_on_dark_kept"] += 1
                kept.append(sp)
            if not kept:
                continue
            lt = "".join(c["c"] for sp in kept for c in sp["chars"])
            if RE_TYPESET_LABEL.search(lt.strip()):
                dropped.append({"reason": "typeset_label", "text": lt.strip(),
                                "bbox": [round(v, 1) for v in l.get("bbox", (0, 0, 0, 0))]})
                stat["dropped_typeset_label"] += 1
                continue
            out.append((l, kept))
    return out, dropped, stat


def classify_pdf_page(pg, spans=None, raster=None):
    """逐页判定文字层可靠性。返回 dict（fine_class / page_type / 指标）。spans 为 page_spans() 结果，可省。"""
    import fitz
    raster = raster or _PageRaster(pg)
    if spans is None:
        spans = page_spans(pg, raster)
    lines, dropped, sstat = spans
    vis, rects, fonts = [], [], collections.Counter()
    subst_lines = 0
    for l, kept in lines:
        lt = "".join(c["c"] for sp in kept for c in sp["chars"])
        if RE_OCR_CIRCLE_SUBST.match(lt):
            subst_lines += 1
        for sp in kept:
            t = "".join(c["c"] for c in sp["chars"])
            vis.append(t)
            k = len(re.sub(r"\s", "", t))
            if k:
                fonts[sp.get("font", "")] += k
                rects.append(sp["bbox"])
    wm = sum(len(d["text"]) for d in dropped)
    raw_s = re.sub(r"\s+", "", "".join(vis))
    s = fix_founder_punct(raw_s)
    mapped = sum(1 for _ in RE_FOUNDER.finditer(raw_s))
    n = len(s)
    inv = tot = 0
    try:
        for t in pg.get_texttrace():
            k = len(t.get("chars", ()))
            tot += k
            if t.get("type") == 3 or t.get("opacity", 1) == 0:
                inv += k
    except Exception:
        pass
    irects = []
    try:
        for info in pg.get_image_info():
            r = fitz.Rect(info["bbox"]) & pg.rect
            if not r.is_empty:
                irects.append(r)
    except Exception:
        pass
    cov = _coverage(irects, pg.rect.width, pg.rect.height)
    try:
        nd = len(pg.get_drawings())
    except Exception:
        nd = -1
    cjk, bad, asym = len(CJK.findall(s)), len(BAD.findall(s)), len(SYM.findall(s))
    invr = inv / float(tot) if tot else 0.0
    cs_chars = sum(v for f, v in fonts.items() if RE_CLEARSCAN_FONT.search(f or ""))
    cs_share = cs_chars / float(sum(fonts.values()) or 1)
    circled = len(re.findall(r"[①②③④]", s))
    ink = None
    if n < 30:
        # 正文（去掉排版标签后）不到 30 字：看墨迹是否在文字层之外（矢量描边字/矢量图）
        if cov >= 0.3:
            cls = "scan_no_text"
        else:
            ink = raster.ink_outside(rects)
            if ink >= 0.012 or nd >= 150:
                cls = "vector_outlined_text"
            elif n > 0:
                cls = "text_reliable"   # 短页：文字层就是全部内容
            else:
                cls = "empty"
    elif invr >= 0.5:
        cls = "scan_with_ocr_layer"
    elif cs_share >= 0.3 or (subst_lines >= 2 and circled == 0):
        cls = "ocr_text_layer"          # ClearScan 等 OCR 文字层：字形与文字层不一一对应
    elif bad / float(n) > 0.05 or (bad >= 5 and bad / float(n) > 0.01) or (n >= 150 and cjk / float(n) < 0.15):
        cls = "garbled_text_layer"
    elif asym / float(n) > 0.03:
        cls = "symbol_garbled"
    elif cov >= 0.85:
        cls = "ocr_layer_over_scan"
    elif cov >= 0.15:
        cls = "text_plus_figure"
    else:
        cls = "text_reliable"
    out = {"fine_class": cls, "page_type": COARSE[cls], "chars": n, "cjk": cjk, "bad": bad, "ascii_sym": asym,
           "founder_punct_mapped": mapped,
           "invisible_ratio": round(invr, 3), "img_cov": round(cov, 3), "drawings": nd, "watermark_chars": wm,
           "clearscan_font_share": round(cs_share, 3), "ocr_circle_subst_lines": subst_lines,
           "dropped_spans": dict((k[8:], v) for k, v in sstat.items() if k.startswith("dropped_")),
           "white_on_dark_kept": sstat.get("white_on_dark_kept", 0)}
    if ink is not None:
        out["ink_outside_text"] = round(ink, 4)
    return out


def _is_dot_span(sp):
    t = "".join(c["c"] for c in sp["chars"]).strip()
    return bool(t) and bool(re.fullmatch(r"[%s]+" % DOTS, t))


def _emphasis_marks(all_lines):
    """着重号：PDF 里是一串“．”单独成 span，比被标字低约半行，常被归到别的行。
    按页找：点的中心 x 落在某个文字字形内、且该字在点的正上方（垂直距离 < 1.2 字高）→ 标记该字，丢掉点。"""
    chars = []
    for spans in all_lines:
        for sp in spans:
            if _is_dot_span(sp):
                continue
            for ci, c in enumerate(sp["chars"]):
                if c["c"].strip():
                    chars.append((id(sp), ci, c["bbox"]))
    marked, dropped = set(), set()
    for spans in all_lines:
        for sp in spans:
            if not _is_dot_span(sp):
                continue
            hits = 0
            for dc in sp["chars"]:
                if not dc["c"].strip():
                    continue
                cx = (dc["bbox"][0] + dc["bbox"][2]) / 2.0
                cy = (dc["bbox"][1] + dc["bbox"][3]) / 2.0
                h = dc["bbox"][3] - dc["bbox"][1]
                best = None
                for sid_, ci, bb in chars:
                    if bb[0] - 0.5 <= cx <= bb[2] + 0.5:
                        ty = (bb[1] + bb[3]) / 2.0
                        dy = cy - ty
                        if 0.5 < dy < 1.2 * max(h, bb[3] - bb[1]) and (best is None or dy < best[0]):
                            best = (dy, sid_, ci)
                if best:
                    marked.add((best[1], best[2]))
                    hits += 1
            if hits and hits >= len([c for c in sp["chars"] if c["c"].strip()]) * 0.6:
                dropped.add(id(sp))
    return marked, dropped


def _line_from_spans(spans, marked=frozenset(), dropped=frozenset()):
    """把一行的 span 合成文字；已判为着重号的点 span 丢掉，被标的字包 <em class="dot">。"""
    spans2 = sorted([sp for sp in spans if id(sp) not in dropped], key=lambda sp: sp["bbox"][0])
    base = max([sp["size"] for sp in spans2] or [10])
    out, n_em = [], 0
    for sp in spans2:
        s, run = "", False
        for ci, c in enumerate(sp["chars"]):
            m = (id(sp), ci) in marked
            if m and not run:
                s += '<em class="dot">'
                run = True
                n_em += 1
            if not m and run:
                s += "</em>"
                run = False
            s += c["c"]
        if run:
            s += "</em>"
        if sp["size"] < base * 0.75 and (sp["flags"] & 1):
            s = "<sup>%s</sup>" % s
        out.append(s)
    return "".join(out), base, n_em


def pdf_lines(pg, spans=None):
    """可见文字行（page_spans 已剔除水印、白底白字、排版标签），按阅读顺序；横版 A3 做双栏检测；着重号转 <em class="dot">。"""
    if spans is None:
        spans = page_spans(pg)
    raw = [(l, sp) for l, sp in spans[0] if sp]
    marked, dropped = _emphasis_marks([sp for _, sp in raw])
    lines = []
    n_em = 0
    n_rad = [0]
    for l, spans in raw:
        t, base, k = _line_from_spans(spans, marked, dropped)
        t2 = fix_radicals(fix_founder_punct(t))
        if t2 != t:
            n_rad[0] += 1
            t = t2
        n_em += k
        if not t.strip():
            continue
        kept = [sp for sp in spans if id(sp) not in dropped]
        x0 = min(sp["bbox"][0] for sp in kept)
        y0 = min(sp["bbox"][1] for sp in kept)
        x1 = max(sp["bbox"][2] for sp in kept)
        y1 = max(sp["bbox"][3] for sp in kept)
        lines.append({"t": t, "x0": x0, "y0": y0, "x1": x1, "y1": y1, "size": base})
    W_ = pg.rect.width
    mid = W_ / 2
    left = sum(1 for l in lines if l["x1"] < mid + 5)
    right = sum(1 for l in lines if l["x0"] > mid - 5)
    two_col = W_ > pg.rect.height and left > 5 and right > 5
    if two_col:
        lines.sort(key=lambda l: (0 if l["x0"] < mid - 5 else 1, round(l["y0"]), l["x0"]))
    else:
        # 同一视觉行的碎片 y0 常差 0.5—1pt（按 y0/3 取整分桶会把后半行排到前半行前面）：先按 y0 聚成行，行内按 x0
        lines.sort(key=lambda l: (l["y0"], l["x0"]))
        rows_, cur_, y_ref = [], [], None
        for l in lines:
            if cur_ and abs(l["y0"] - y_ref) < max(3.0, 0.4 * l["size"]):
                cur_.append(l)
            else:
                if cur_:
                    rows_.append(cur_)
                cur_, y_ref = [l], l["y0"]
        if cur_:
            rows_.append(cur_)
        lines = [l for r_ in rows_ for l in sorted(r_, key=lambda z: z["x0"])]
        merged = []
        for l in lines:
            if merged and abs(merged[-1]["y0"] - l["y0"]) < 3 and l["x0"] >= merged[-1]["x1"] - 2:
                gap = l["x0"] - merged[-1]["x1"]
                merged[-1]["t"] += ("\t" if gap > merged[-1]["size"] * 2 else "") + l["t"]
                merged[-1]["x1"] = l["x1"]
            else:
                merged.append(dict(l))
        lines = merged
    pdf_lines.last_radicals = n_rad[0]
    return lines, two_col, n_em


def _label_len(pt):
    m = re.match(r"^\s*(\d{1,2}\s*[．.、]|[A-DＡ-Ｄ]\s*[．.、]|[（(]\s*\d{1,2}\s*[）)]|[①-⑳])", pt)
    if not m:
        return 0.0
    return sum(1.0 if ord(c) > 0x2000 else 0.55 for c in m.group(0))


def lines_to_paras(lines, left_margin, right_margin):
    """行 → 段。规则：新题号/选项/小问/材料/圈码起首必断；行距过大必断；上一行排满则续行
    （上一行句末且本行首行缩进时才断）；上一行没排满且以句末标点结束则断；悬挂缩进（题号后对齐）算续行。"""
    paras = []
    for l in lines:
        t = l["t"].strip()
        pt = plain(t)
        if not paras:
            paras.append(dict(l, t=t, hang=l["x0"] + _label_len(pt) * l["size"]))
            continue
        prev = paras[-1]
        label = bool(RE_QSTART.match(pt) or RE_OPT.match(pt) or RE_SUBQ.match(pt) or RE_MAT.match(pt) or
                     RE_CIRC.match(pt) or RE_ANS_INLINE.match(pt))
        indent = l["x0"] - left_margin > l["size"] * 1.2
        hanging = abs(l["x0"] - prev.get("hang", -99)) < l["size"] * 0.9 and prev.get("hang", 0) > left_margin + 1
        prev_full = prev["x1"] >= right_margin - prev["size"] * 2.0
        prev_end = plain(prev["t"]).rstrip()[-1:] in TERMINAL
        gap = l["y0"] - prev["y1"] > l["size"] * 1.2 or l["y0"] < prev["y0"] - 2
        if label or gap or "\t" in prev["t"]:
            new = True
        elif prev_full:
            new = prev_end and indent and not hanging
        elif prev_end:
            new = True
        else:
            new = indent and not hanging
        if new:
            paras.append(dict(l, t=t, hang=l["x0"] + _label_len(pt) * l["size"]))
        else:
            sep = " " if re.match(r"[A-Za-z0-9]", pt[:1] or "") and re.search(r"[A-Za-z0-9]$", plain(prev["t"])) else ""
            prev["t"] = prev["t"].rstrip() + sep + t
            prev["x1"] = max(prev["x1"], l["x1"])
            prev["y1"] = l["y1"]
    return paras


def _pct(vals, q):
    v = sorted(vals)
    if not v:
        return 0.0
    return v[min(len(v) - 1, max(0, int(round(q * (len(v) - 1)))))]


def split_options(t):
    """“A．①② B．①③ C．②④ D．③④” 拆成四行；只在确有 2 个以上选项字母时拆。"""
    pt = plain(t)
    if len(re.findall(r"(?:^|\s)[A-DＡ-Ｄ]\s*[．.、]", pt)) >= 2 and RE_OPT.match(pt):
        return [x.strip() for x in re.split(r"[\s\t\u3000]+(?=(?:<[^>]+>|\*\*)*[B-DＢ-Ｄ]\s*[．.、])", t) if x.strip()]
    return [t]


def ocr_paras(txt):
    out = []
    for raw in txt.splitlines():
        t = raw.strip()
        if not t or RE_PAGENO.match(t):
            continue
        if out and not (RE_QSTART.match(t) or RE_OPT.match(t) or RE_SUBQ.match(t) or RE_MAT.match(t) or
                        RE_CIRC.match(t)) and out[-1][-1:] not in TERMINAL and len(out[-1]) > 12:
            out[-1] += t
        else:
            out.append(t)
    return out


def find_ocr_text(sid, page, ocr_dirs):
    for d in ocr_dirs:
        if not d:
            continue
        for name in ("p%03d.ocr.txt" % page, "p%d.ocr.txt" % page, "page_%03d.txt" % page):
            f = os.path.join(d, name)
            if os.path.isfile(f):
                txt = open(f, encoding="utf-8", errors="replace").read()
                return re.sub(r"<<<OCR [^>]*>>>", "", txt), f
    return None, None


def run_ocr(path, pages, out_dir):
    """可选：用 macOS Vision（ocr-vision）只对缺 OCR 的页补跑，产物写 out_dir。"""
    exe = ocr_exe()
    if not exe or not pages:
        return {}
    os.makedirs(out_dir, exist_ok=True)
    spec = ",".join(str(p) for p in pages)
    raw = os.path.join(out_dir, "_raw")
    os.makedirs(raw, exist_ok=True)
    try:
        subprocess.run([exe, "--out", raw, "--pages", spec, path], check=False, capture_output=True,
                       timeout=max(600, 60 + 15 * len(pages)))
    except Exception:
        return {}
    # ocr-vision 把多页写进同一个文件（名如 X.pages_2_4-5.ocr.txt），页头“===== X.pdf page N =====”。
    # 旧写法按文件名末尾数字取页码：多页时只认出最后一页、且那一页拿到全部页的文字（2026-09-24 修）。现按页头拆成逐页文件。
    got = {}
    for f in sorted(os.listdir(raw)):
        if not f.endswith(".txt"):
            continue
        cur, buf = None, []
        for l in open(os.path.join(raw, f), encoding="utf-8", errors="replace").read().splitlines() + \
                ["===== end page 0 ====="]:
            m = re.match(r"^=====\s*.+?\s+page\s+(\d+)\s*=====$", l.strip())
            if m:
                if cur is not None and cur in pages:
                    fp = os.path.join(out_dir, "p%03d.ocr.txt" % cur)
                    with open(fp, "w", encoding="utf-8") as fh:
                        fh.write("\n".join(buf))
                    got[cur] = fp
                cur, buf = int(m.group(1)), []
            else:
                buf.append(l)
    shutil.rmtree(raw, ignore_errors=True)
    return got


def extract_pdf(path, sid, asset_root=None, ocr_dirs=(), render_dpi=110, ocr_run_dir=None):
    import fitz
    try:
        doc = fitz.open(path)
    except Exception as e:
        raise RuntimeError("PDF 打不开（%s）" % type(e).__name__)
    if doc.needs_pass or doc.is_encrypted:
        raise RuntimeError("PDF 已加密，脚本不解密，请人工处理")
    if len(doc) == 0:
        raise RuntimeError("PDF 没有页面")
    blocks, pages = [], []
    flags = collections.Counter()
    hf = collections.Counter()
    page_lines, metas = [], []
    for pg in doc:
        raster = _PageRaster(pg)
        spans = page_spans(pg, raster)
        ls, two_col, n_em = pdf_lines(pg, spans)
        page_lines.append((ls, two_col))
        flags["emphasis_dot"] += n_em
        flags["codepoint_fixed_lines"] += getattr(pdf_lines, "last_radicals", 0)
        H = pg.rect.height
        for l in ls:
            if l["y0"] < H * 0.07 or l["y1"] > H * 0.93:
                hf[re.sub(r"\d+", "#", re.sub(r"\s+", "", plain(l["t"])))] += 1
        meta = classify_pdf_page(pg, spans, raster)
        meta["dropped"] = spans[1]
        for d_ in spans[1]:
            flags["dropped_" + d_["reason"]] += 1
        if spans[2].get("white_on_dark_kept"):
            flags["white_on_dark_kept"] += spans[2]["white_on_dark_kept"]
        metas.append(meta)
    npg = len(doc)
    hf_keys = {k for k, v in hf.items() if v >= max(2, npg * 0.5)}
    need_ocr = [i + 1 for i, m in enumerate(metas) if m["fine_class"] in ("scan_no_text", "vector_outlined_text")
                and not find_ocr_text(sid, i + 1, ocr_dirs)[0]]
    extra_ocr = run_ocr(path, need_ocr, ocr_run_dir) if ocr_run_dir else {}
    # OCR 文本里跨页重复的短行（扫描 App 水印、页眉页脚）只记一次、不进正文
    ocr_rep = collections.Counter()
    for i in range(1, npg + 1):
        txt, _ = find_ocr_text(sid, i, ocr_dirs)
        if not txt and i in extra_ocr:
            txt = open(extra_ocr[i], encoding="utf-8", errors="replace").read()
        for ln in set(x.strip() for x in (txt or "").splitlines() if 0 < len(x.strip()) <= 24):
            ocr_rep[ln] += 1
    ocr_drop = {k for k, v in ocr_rep.items() if v >= max(3, npg * 0.4)}
    if ocr_drop:
        flags["ocr_repeated_lines_dropped"] += len(ocr_drop)
    for i, pg in enumerate(doc, 1):
        meta = dict(metas[i - 1], page=i)
        cls = meta["fine_class"]
        blocks.append({"kind": "pagemark", "text": "", "page": i, "class": cls})
        H = pg.rect.height
        ls = [l for l in page_lines[i - 1][0]
              if not ((l["y0"] < H * 0.07 or l["y1"] > H * 0.93) and
                      (re.sub(r"\d+", "#", re.sub(r"\s+", "", plain(l["t"]))) in hf_keys or
                       RE_PAGENO.match(plain(l["t"]).strip())))]
        # 页边侧标（如左缘黑底白字“政”）：整行落在页宽两侧 6% 内的短字，不是题面；剔除并记入 dropped
        Wp = pg.rect.width
        keep_ls = []
        for l in ls:
            if (l["x1"] < Wp * 0.06 or l["x0"] > Wp * 0.94) and len(re.sub(r"\s", "", plain(l["t"]))) <= 4:
                meta.setdefault("dropped", []).append({"reason": "margin_tab", "text": plain(l["t"]).strip(),
                                                       "bbox": [round(l["x0"], 1), round(l["y0"], 1),
                                                                round(l["x1"], 1), round(l["y1"], 1)]})
                flags["dropped_margin_tab"] += 1
                continue
            keep_ls.append(l)
        ls = keep_ls
        meta["two_column"] = page_lines[i - 1][1]
        if cls in ("scan_no_text", "vector_outlined_text", "empty"):
            txt, src = find_ocr_text(sid, i, ocr_dirs)
            if not txt and i in extra_ocr:
                txt, src = open(extra_ocr[i], encoding="utf-8", errors="replace").read(), extra_ocr[i]
            if txt:
                txt = "\n".join(x for x in txt.splitlines() if x.strip() not in ocr_drop)
                for t in ocr_paras(txt):
                    for s in split_options(t):
                        blocks.append({"kind": "p", "text": s, "page": i, "ocr": True})
                meta["ocr_source"] = src
                flags["ocr_pages"] += 1
            elif cls != "empty":
                flags["no_text_pages"] += 1
            if cls != "empty":
                rel = "assets/%s/p%03d.png" % (sid, i)
                if asset_root and render_dpi:
                    full = os.path.join(asset_root, rel)
                    os.makedirs(os.path.dirname(full), exist_ok=True)
                    if not os.path.isfile(full):
                        pg.get_pixmap(dpi=render_dpi).save(full)
                blocks.append({"kind": "scanpage", "text": "", "path": rel, "page": i, "class": cls})
            pages.append(meta)
            continue
        ocr_like = cls in OCR_LAYER_CLASSES
        garbled = cls in ("symbol_garbled", "garbled_text_layer")
        tabs = []
        try:
            for tb in pg.find_tables().tables:
                if tb.row_count >= 2 and tb.col_count >= 2:
                    data = tb.extract()
                    rects = [fitz.Rect(r.bbox) for r in tb.rows]
                    # find_tables 偶尔把表下方的下一题题号行吞进表格：在题号行处截断
                    cut_y = None
                    for l in ls:
                        c = fitz.Rect(l["x0"], l["y0"], l["x1"], l["y1"])
                        # 行框常比表格框宽（整行排满），按重叠面积判断是否在表内
                        if not (_rarea(fitz.Rect(tb.bbox) & c) > 0.5 * _rarea(c) and c.y0 > rects[0].y1 - 1):
                            continue
                        pt_ = plain(l["t"])
                        full_w = tb.col_count >= 2 and c.width >= 0.8 * fitz.Rect(tb.bbox).width
                        if re.match(r"^\s*\d{1,2}\s*[．.、]\s*(?:[（(]\s*\d{1,2}\s*分|\S{6,})", pt_) or \
                                (full_w and re.match(r"^\d{1,2}[．.、](?:[（(]\d{1,2}分|\S{6,})", re.sub(r"\s+", "", pt_))):
                            # 多栏表里横跨整表宽度的“N．…”行不可能是单元格内容，是被吞进来的下一题
                            cut_y = c.y0 if cut_y is None else min(cut_y, c.y0)
                    if cut_y is not None:
                        keep = [i for i, r in enumerate(rects) if r.y1 <= cut_y + 1]
                        flags["table_cut_at_question_line"] += 1
                        if len(keep) < 2:
                            continue
                        data = [data[i] for i in keep]
                        bb = fitz.Rect(rects[keep[0]])
                        for i in keep:
                            bb |= rects[i]
                    else:
                        bb = fitz.Rect(tb.bbox)
                    rows = [[{"text": squeeze((c or "").replace("\n", "")), "span": 1, "vm": None} for c in r]
                            for r in data]
                    empty = sum(1 for r in rows for c in r if not c["text"].strip()) / float(sum(len(r) for r in rows) or 1)
                    multi = sum(1 for r in rows if sum(1 for c in r if c["text"].strip()) >= 2) / float(len(rows) or 1)
                    if len(rows) >= 3 and multi < 0.3:
                        # 几乎每行只有一个非空格子：是装饰框线围住的正文（对话框、材料框），不是表格；按正文段落处理
                        flags["pseudo_table_as_text"] += 1
                        continue
                    tabs.append((bb, rows, empty))
        except Exception:
            pass
        figs = []
        if not ocr_like:
            for info in pg.get_image_info():
                r = fitz.Rect(info["bbox"]) & pg.rect
                if r.width * r.height > pg.rect.width * pg.rect.height * 0.01:
                    figs.append(("image", r))
            try:
                for r in pg.cluster_drawings():
                    if r.width * r.height > pg.rect.width * pg.rect.height * 0.02 and \
                            not any(r.intersects(t[0]) for t in tabs):
                        figs.append(("vector", r))
            except Exception:
                pass
        items = []
        for l in ls:
            c = fitz.Rect(l["x0"], l["y0"], l["x1"], l["y1"])
            if any(t[0].contains(c) or _rarea(t[0] & c) > 0.6 * _rarea(c) for t in tabs):
                continue
            items.append(l)
        lm = _pct([l["x0"] for l in items], 0.1) if items else 0
        rm = _pct([l["x1"] for l in items], 0.9) if items else pg.rect.width
        paras = lines_to_paras(items, lm, rm)
        seq = [("p", p["y0"], p) for p in paras] + [("t", t[0].y0, t) for t in tabs] + [("f", f[1].y0, f) for f in figs]
        if not meta["two_column"]:
            seq.sort(key=lambda x: x[1])
        fi = 0
        for kind, y, obj in seq:
            if kind == "p":
                for s in [squeeze(x) for x in split_options(obj["t"])]:
                    blocks.append({"kind": "p", "text": s, "page": i, "ocr": ocr_like, "garbled": garbled,
                                   "bbox": [round(obj["x0"], 1), round(obj["y0"], 1), round(obj["x1"], 1),
                                            round(obj["y1"], 1)]})
            elif kind == "t":
                rows = obj[1]
                blocks.append({"kind": "table", "text": table_md(rows, False), "page": i, "rows": len(rows),
                               "cols": table_shape(rows), "merged": obj[2] > 0.5, "empty_ratio": round(obj[2], 2),
                               "bbox": [round(v, 1) for v in obj[0]]})
                flags["table"] += 1
            else:
                fi += 1
                rel = "assets/%s/p%03d_fig%02d.png" % (sid, i, fi)
                if asset_root:
                    full = os.path.join(asset_root, rel)
                    os.makedirs(os.path.dirname(full), exist_ok=True)
                    pg.get_pixmap(dpi=150, clip=obj[1]).save(full)
                blocks.append({"kind": "image", "text": "", "path": rel, "page": i, "what": obj[0],
                               "bbox": [round(v, 1) for v in obj[1]]})
                flags["figure_" + obj[0]] += 1
        if ocr_like:
            flags["ocr_layer_pages"] += 1
            rel = "assets/%s/p%03d.png" % (sid, i)
            if asset_root and render_dpi:
                full = os.path.join(asset_root, rel)
                os.makedirs(os.path.dirname(full), exist_ok=True)
                if not os.path.isfile(full):
                    pg.get_pixmap(dpi=render_dpi).save(full)
            blocks.append({"kind": "scanpage", "text": "", "path": rel, "page": i, "class": cls})
        if garbled:
            flags["garbled_pages"] += 1
        pages.append(meta)
    return blocks, flags, pages


def _rarea(r):
    return max(0.0, r.width) * max(0.0, r.height) if not r.is_empty else 0.0


# ---------------------------------------------------------------- 旧格式
def legacy_to_docx(path, work_dir):
    """doc/rtf → docx：macOS 自带 textutil（确定性、离线）；失败返回 None。"""
    exe = shutil.which("textutil") or "/usr/bin/textutil"
    if not os.path.isfile(exe):
        return None
    os.makedirs(work_dir, exist_ok=True)
    out = os.path.join(work_dir, os.path.splitext(os.path.basename(path))[0] + ".converted.docx")
    try:
        subprocess.run([exe, "-convert", "docx", path, "-output", out], check=True, capture_output=True, timeout=120)
    except Exception:
        return None
    return out if os.path.isfile(out) else None


# ---------------------------------------------------------------- 统一入口：抽取
def _docx_blocks(path, sid, asset_root):
    try:
        d = Docx(path, sid, asset_root)
        if d._xml("word/document.xml") is None:
            raise RuntimeError("DOCX 缺少 word/document.xml")
        blocks = d.extract()
    except zipfile.BadZipFile:
        raise RuntimeError("DOCX 不是有效的 zip 包（空文件或已损坏）")
    except ET.ParseError as e:
        raise RuntimeError("DOCX 的 XML 损坏（%s）" % e)
    return blocks, d.flags, docx_page_class(blocks)


def detect_format(path):
    ext = os.path.splitext(path)[1].lower().lstrip(".")
    return {"docx": "docx", "pdf": "pdf", "pptx": "pptx", "doc": "doc", "rtf": "rtf"}.get(ext, ext)


def extract_source(path, fmt=None, sid=None, asset_root=None, ocr_dirs=None, render_dpi=110, ocr_run_dir=None,
                   work_dir=None):
    """抽取单个原件 → (blocks, flags, pages, info)。asset_root=None 时不落盘任何图片（体检用）。"""
    t0 = time.time()
    fmt = fmt or detect_format(path)
    digest = sha256(path)
    sid = sid or sid_of(path, digest)
    if ocr_dirs is None:
        ocr_dirs = [os.path.join(BANK, "evidence", sid, "visual_transcript")]
    info = {"source_id": sid, "path": path, "format": fmt, "sha256": digest}
    if fmt == "docx":
        blocks, flags, pages = _docx_blocks(path, sid, asset_root)
    elif fmt == "pptx":
        blocks, flags, pages = extract_pptx(path, sid, asset_root)
    elif fmt == "pdf":
        blocks, flags, pages = extract_pdf(path, sid, asset_root, ocr_dirs, render_dpi, ocr_run_dir)
    elif fmt in ("doc", "rtf"):
        # 转换副本只放调用方给的 work_dir（在 --out-dir 内）；没给时用自动清理的临时目录，不在系统临时区留原件副本
        tmp = None
        if not work_dir:
            tmp = tempfile.TemporaryDirectory(prefix="doc2md_legacy_")
            work_dir = tmp.name
        try:
            conv = legacy_to_docx(path, work_dir)
            if not conv:
                raise RuntimeError("旧格式 %s 转换失败（textutil 不可用或报错）" % fmt)
            blocks, flags, pages = _docx_blocks(conv, sid, asset_root)
        finally:
            if tmp is not None:
                tmp.cleanup()
        flags["legacy_converted_textutil"] += 1
        pages = [dict(p, fine_class="legacy_converted", page_type=COARSE["legacy_converted"]) for p in pages]
        info["converted_from"] = fmt
    else:
        raise RuntimeError("不支持的格式：%s" % fmt)
    cls_of = {p["page"]: p["fine_class"] for p in pages}
    for b in blocks:
        if "page" in b and "class" not in b:
            b["class"] = cls_of.get(b["page"], "docx" if fmt == "docx" else "?")
    info["seconds"] = round(time.time() - t0, 3)
    info["flags"] = dict(flags)
    info["page_types"] = dict(collections.Counter(p["page_type"] for p in pages))
    info["fine_classes"] = dict(collections.Counter(p["fine_class"] for p in pages))
    return blocks, flags, pages, info


# ---------------------------------------------------------------- 切题
RE_ANS_HEAD = re.compile(r"^(.{0,40}(参考答案|答案及评分|评分参考|答案与解析|试题答案|答案和评分)[^。]{0,20})$")
RE_ANS_INLINE = re.compile(r"^\s*(\d{1,2}\s*[．.、]?\s*)?【\s*(答案|解析|详解|分析|点睛|点评|考点|名师点睛|试题解析|考点定位|答案解析|详析)\s*】")
RE_KEYLINE = re.compile(r"^\s*(?:【\s*答案\s*】)?\s*(\d{1,2}\s*[．.、:：]?\s*[A-D](?![A-Za-z])[\s\t|，,;；]*){2,}")
RE_SECTION = re.compile(r"(第[一二三四]部分)|(本部分共\s*\d+\s*(小)?题)")
RE_SEC_EACH = re.compile(r"本部分共\s*(\d+)\s*(?:小)?题[，,、]?\s*每(?:小)?题\s*(\d+)\s*分[，,、]?\s*共\s*(\d+)\s*分")
RE_SEC_TOTAL = re.compile(r"本部分共\s*(\d+)\s*(?:小)?题[，,、]?\s*共\s*(\d+)\s*分")
RE_SCORE = re.compile(r"[（(]\s*(?:共\s*)?(\d{1,2})\s*分\s*[）)]")


def _qnum(t, ocr=False, expected=None):
    m = RE_QSTART.match(t)
    if m:
        return int(m.group(1))
    if ocr and expected is not None:
        # OCR 常把题号后的“．”认成“：/，/-/空格”，或在前面多出“、”；只在恰为下一题号时接受
        m2 = re.match(r"^\s*[、，,.．·]?\s*(\d{1,2})\s*(?:[：:，,\-—_]|\s+(?=\S))", t)
        if m2 and int(m2.group(1)) == expected:
            return expected
        m3 = re.match(r"^\s*[、，,.．·]\s*(\d{1,2})\s*[．.、]", t)
        if m3 and int(m3.group(1)) == expected:
            return expected
    return None


def split_paper(blocks):
    """原卷/教师版：按 1..N 期望号切题；识别卷内答案区（号码重置或“参考答案”标题）与内附【答案】【解析】。

    返回 dict(front, questions{n:{stem,answer,pages}}, answer_region, sections, events)
    """
    qs = collections.OrderedDict()
    front, answer_region, sections, events = [], [], [], []
    cur, expected, in_ans_region, inline_ans = None, 1, False, False
    last_page = None
    for b in blocks:
        if b["kind"] == "pagemark":
            last_page = b["page"]
            (answer_region if in_ans_region else (qs[cur]["stem"] if cur is not None else front)).append(b)
            continue
        t = plain(b.get("text", "")).strip()
        if b["kind"] == "table":
            t = re.sub(r"^\|\s*", "", t.split("\n")[0]) if t.startswith("|") else re.sub(r"<[^>]+>", " ", t).strip()
        if in_ans_region:
            answer_region.append(b)
            continue
        if b["kind"] in ("p", "textbox", "table") and t:
            sec = RE_SEC_EACH.search(t) or RE_SEC_TOTAL.search(t)
            if RE_SECTION.search(t):
                seen_same = any(s["text"][:12] == t[:12] for s in sections)
                if cur is not None and cur >= 5 and seen_same:
                    in_ans_region = True
                    events.append({"event": "answer_region_by_repeated_section_header", "page": b.get("page"),
                                   "text": clip(t, 40)})
                    answer_region.append(b)
                    continue
                if sec:
                    g = sec.groups()
                    sections.append({"text": clip(t, 80), "start_q": expected, "count": int(g[0]),
                                     "per_q": int(g[1]) if len(g) == 3 else None, "total": int(g[-1]),
                                     "choice": len(g) == 3 or ("选择" in t and "非选择" not in t), "page": b.get("page")})
                else:
                    sections.append({"text": clip(t, 80), "start_q": expected, "count": None, "per_q": None,
                                     "total": None, "choice": "选择" in t and "非选择" not in t, "page": b.get("page")})
            if cur is not None and cur >= 5 and RE_ANS_HEAD.match(t) and len(t) <= 60 and not RE_QSTART.match(t):
                in_ans_region = True
                events.append({"event": "answer_region_by_heading", "page": b.get("page"), "text": clip(t, 40)})
                answer_region.append(b)
                continue
            n = _qnum(t, b.get("ocr"), expected)
            if n is not None:
                if cur is not None and cur >= 5 and n == 1 and (RE_KEYLINE.match(t) or RE_ANS_INLINE.match(t) or
                                                               re.match(r"^\s*1\s*[．.、]\s*[A-D](\s|$)", t)):
                    in_ans_region = True
                    events.append({"event": "answer_region_by_number_reset", "page": b.get("page"), "text": clip(t, 40)})
                    answer_region.append(b)
                    continue
                ok = n == expected or (cur is not None and n in (expected + 1, expected + 2) and len(t) > 8)
                if ok and RE_KEYLINE.match(t):
                    ok = False
                if ok:
                    if n != expected:
                        events.append({"event": "gap", "expected": expected, "got": n, "page": b.get("page")})
                    cur, expected, inline_ans = n, n + 1, False
                    qs[n] = {"stem": [], "answer": [], "pages": set()}
            if cur is not None and RE_ANS_INLINE.match(t):
                inline_ans = True
        if cur is None:
            front.append(b)
            continue
        qs[cur]["answer" if inline_ans else "stem"].append(b)
        if b.get("page"):
            qs[cur]["pages"].add(b["page"])
    return {"front": front, "questions": qs, "answer_region": answer_region, "sections": sections, "events": events}


RE_R_STRONG = [
    re.compile(r"^\s*第\s*(\d{1,2})\s*题"),
    re.compile(r"^\s*(\d{1,2})\s*[．.、]?\s*题\s*(?:[：:．.、\s（(]|$)"),   # “18题”“22.题”（OCR 常见）也算
    re.compile(r"^\s*【\s*(\d{1,2})\s*题?\s*】"),
    re.compile(r"^\s*(\d{1,2})\s*[．.、]?\s*[（(]\s*(?:共\s*)?\d{1,2}\s*分"),
    re.compile(r"^\s*(\d{1,2})\s*[．.、]?\s*【\s*(?:答案|解析|详解)"),
    re.compile(r"^\s*(\d{1,2})\s*[．.、]?\s*[（(]\s*[1-9一二三四]\s*[）)]"),
    re.compile(r"^\s*(\d{1,2})\s*[．.]\s*[1-9]\s*[．.、]"),
    # “20.评分细则第一问”“21．参考答案”：题号后紧跟评分/答案字样，是题首不是列表项
    re.compile(r"^\s*(\d{1,2})\s*[．.、]?\s*题?\s*(?:评分细则|评分标准|评分参考|阅卷细则|阅卷标准|细则|参考答案|答案要点)"),
]
RE_R_WEAK = re.compile(r"^\s*(\d{1,2})\s*[．.](?!\d)")
RE_R_WEAKEST = re.compile(r"^\s*(\d{1,2})\s*、")
# 答案区题首“5．【分析】/【答案】/【解析】…”：一律另起一块，不受期望序号和题面相似度影响
ANS_HEAD_WORDS = r"(?:分析|答案|解析|详解|解答|点睛|考点|点评|试题解析|名师点睛|考点定位|答案解析|详析)"
RE_ANS_HEAD_N = re.compile(r"^\s*(?:〔OCR候选〕)?(\d{1,2})\s*[．.、]?\s*【\s*%s\s*】" % ANS_HEAD_WORDS)
RE_ANS_HEAD_N_ML = re.compile(r"(?m)^\s*(?:〔OCR候选〕)?(\d{1,2})\s*[．.、]?\s*【\s*%s\s*】" % ANS_HEAD_WORDS)


def _rubric_num(t):
    m = RE_ANS_HEAD_N.match(t)
    if m:
        return int(m.group(1)), "answer_head"
    for rx in RE_R_STRONG:
        m = rx.match(t)
        if m:
            return int(m.group(1)), "strong"
    m = RE_R_WEAK.match(t)
    if m:
        return int(m.group(1)), "weak"
    m = RE_R_WEAKEST.match(t)
    if m:
        return int(m.group(1)), "weakest"
    return None, None


def _grams(s, k=4):
    return {s[i:i + k] for i in range(max(0, len(s) - k + 1))}


def _sim(head, stem):
    """候选题起首文字与原卷该题题面的 4-gram 重合度（0..1）。"""
    a = canon(norm_display(head))[:48]
    g = _grams(a)
    if not g or not stem:
        return 0.0
    return len(g & stem) / float(len(g))


RE_ASK = re.compile(r"结合材料|运用|说明|分析|谈谈|阐述|阐释|指出|概括|为什么|如何|怎样|完成|评析|论证|提出")
RE_DECK_RESTART = re.compile(r"^\s*(一\s*[、．.]|【?试题分析】?|知识板块)")


def split_rubric(blocks, known=None, hint=None, stems=None, subjective=None, answer_region=False, choice=None):
    """评分材料/参考答案/讲评：按题号起首切块。

    known=已知题号集合（整卷模式由原卷给出）；hint=文件名里的题号（分题切片）；
    stems={题号: 原卷题面4-gram集合}，用来区分“16．结合材料…”（复述设问）与细则里的列表项“1、物质决定意识…”。
    做法：列出所有题号起首候选并打分（强式3/弱式1.5/顿号式0.5，复述设问加分，像列表项扣分），
    取题号严格递增、总分最高的一条链作为切点。"""
    cands = []
    for i, b in enumerate(blocks):
        if b["kind"] not in ("p", "textbox", "table"):
            continue
        t = plain(b.get("text", "")).strip()
        if b["kind"] == "table":
            t = re.sub(r"^\|\s*", "", t)
        if RE_KEYLINE.match(t):
            continue
        n, s = _rubric_num(t)
        if n is None:
            continue
        cands.append((i, n, s, t))
    if known is None:
        strong = sorted({n for _, n, s, _ in cands if s in ("strong", "answer_head")})
        weak = sorted({n for _, n, s, _ in cands if s == "weak"})
        if strong:
            lo, hi = min(strong), max(strong)
            known = set(strong) | {n for n in weak if lo <= n <= hi + 1}
        elif weak:
            known = set(weak)
        else:
            known = {n for _, n, _, _ in cands}
    known = set(known or ())
    # 没写题号、但复述了某题设问的段落（如“(1)结合材料，谈谈……(5分)”）：与原卷该题题面 4-gram 重合≥0.6 记为候选起点
    if stems:
        have = {i for i, _, _, _ in cands}
        for i, b in enumerate(blocks):
            if i in have or b["kind"] not in ("p", "textbox"):
                continue
            t = plain(b.get("text", "")).strip()
            if len(t) < 12 or RE_KEYLINE.match(t):
                continue
            if not RE_ASK.search(t[:80]):
                continue  # 只认复述设问的句子，不认复述材料的答案句
            pool = [n for n in stems if n in known and (not subjective or n in subjective)]
            best = max(((_sim(t, stems[n]), n) for n in pool), default=(0, None))
            if best[0] >= 0.6:
                cands.append((i, best[1], "restate", t))
        cands.sort(key=lambda x: x[0])
    scored = []
    for i, n, s, t in cands:
        if n not in known:
            continue
        # 评分材料里“1.社会主义代替资本主义…”这类列表项不能当选择题起点；选择题只认强题首或“1．B”式答案行
        if choice and n in choice and s in ("weak", "weakest", "restate") and \
                not re.match(r"^\s*\d{1,2}\s*[．.、]\s*[A-D](?![A-Za-z])", t):
            continue
        sc = {"answer_head": 6.0, "strong": 3.0, "weak": 1.5, "weakest": 0.5, "restate": 1.0}[s]
        if stems and n in stems and s != "answer_head":
            sim = _sim(re.sub(r"^\s*(第\s*)?\d{1,2}\s*(题)?[．.、]?", "", t), stems[n])
            sc += 3.0 * sim
            # 卷内答案区的“16．①文明因交流…”不复述设问，相似度低是常态，不扣分；评分材料里的列表项才扣
            if sim < 0.15 and s != "strong" and not (answer_region and s == "weak"):
                sc -= 1.5
        if sc > 0:
            scored.append((i, n, s, sc))
    # 最大权严格递增链（O(n^2)，候选通常 <200）
    best = [0.0] * len(scored)
    prev = [-1] * len(scored)
    for j, (i, n, s, sc) in enumerate(scored):
        best[j] = sc
        for h in range(j):
            if scored[h][1] < n and best[h] + sc > best[j]:
                best[j] = best[h] + sc
                prev[j] = h
    chain = []
    if scored:
        j = max(range(len(scored)), key=lambda x: best[x])
        while j >= 0:
            chain.append(scored[j])
            j = prev[j]
        chain.reverse()
    # 同一题号有多个候选时，起点前移到上一题起点之后最早的那个（题首“16（8分）”不丢给上一题）
    fixed, prev_i = [], -1
    for i, n, s_, sc in chain:
        # “N．【分析】”题首本身就是起点，不再前移（前面同号的“2．文化与经济…”是上一题解析里的列表项）
        earliest = i if s_ == "answer_head" else \
            min((ii for ii, nn, _, _ in scored if nn == n and prev_i < ii <= i), default=i)
        fixed.append((earliest, n, s_, sc))
        prev_i = earliest
    chain = fixed
    starts = {i: n for i, n, s, sc in chain}
    events = []
    # 讲评/评分 PPT 常见结构“一、试题分析 二、评分细则 三、典型示例 四、学生问题 五、复练试题”，每题一轮、常不写题号。
    # 每页首段是“一、/试题分析/知识板块”的算一轮起点；有显式题号的轮次用题号，没有的按主观题顺序补
    # （补出来的号记 split_uncertain）。显式题号若落在某一轮内部，只当旁证，不再另起一题。
    if subjective:
        subj = sorted(x for x in subjective if x in known or not known)
        restarts, seen_page = [], set()
        for i, b in enumerate(blocks):
            if b["kind"] in ("p", "textbox") and b.get("page") not in seen_page and plain(b.get("text", "")).strip():
                seen_page.add(b.get("page"))
                if RE_DECK_RESTART.match(plain(b["text"]).strip()):
                    restarts.append(i)
        if len(restarts) >= 2:
            anchors = sorted((i, n) for i, n in starts.items() if n in subj)
            bounds = sorted(set(restarts) | {i for i, _ in anchors})
            new_starts, prev, inferred = {}, None, []
            for k_, i in enumerate(bounds):
                nxt_b = bounds[k_ + 1] if k_ + 1 < len(bounds) else len(blocks)
                inside = [n for j, n in anchors if i <= j < nxt_b and (prev is None or n > prev)]
                if i in starts and starts[i] in subj and (prev is None or starts[i] > prev):
                    n = starts[i]
                elif inside:
                    n = inside[0]
                else:
                    later = [n for j, n in anchors if j >= nxt_b]
                    cand_n = [x for x in subj if (prev is None or x > prev) and (not later or x < later[0])]
                    if not cand_n:
                        continue
                    n = cand_n[0]
                    inferred.append(n)
                if prev is not None and n <= prev:
                    continue
                new_starts[i] = n
                prev = n
            if new_starts:
                starts = new_starts
                chain = [(i, n, "restart" if n in inferred else "anchor", 0) for i, n in sorted(starts.items())]
                events.append({"event": "rubric_sections_by_restart", "q": sorted(starts.values()),
                               "inferred": inferred, "strength": "restart"})

    qs = collections.OrderedDict()
    pre = []
    cur = hint if hint is not None and (not known or hint in known) else None
    if cur is not None:
        # 文件名题号与首个锚点同号而锚点前还有标题块时，先建键（原先这里 KeyError）
        qs.setdefault(cur, [])
        if not chain or chain[0][1] != cur:
            events.append({"event": "rubric_start_by_filename", "q": cur})
    for i, b in enumerate(blocks):
        if i in starts:
            cur = starts[i]
            qs.setdefault(cur, [])
            events.append({"event": "rubric_start", "q": cur, "page": b.get("page"),
                           "strength": next((s_ for ii, n, s_, sc in chain if ii == i), "?")})
        (qs[cur] if cur is not None else pre).append(b)
    return pre, qs, events


def answer_key(blocks, qnum=None):
    """选择题答案键：行内“1．B 2．D…”、区间“1-5 BCADB”、题号/答案表、单题“【答案】B”。返回 {题号:(字母,页)}
    qnum 给定时（已切好的单题块），“【答案】B / 答案：B”记到该题。"""
    key = {}
    for b in blocks:
        t = plain(b.get("text", ""))
        pg = b.get("page")
        if qnum is not None and b["kind"] in ("p", "textbox"):
            m = re.match(r"^\s*(?:\d{1,2}\s*[．.、]?\s*)?(?:【\s*答案\s*】|答案\s*[：:])\s*([A-D])(?![A-Za-z])", t)
            if m and qnum not in key:
                key[qnum] = (m.group(1), pg)
            m = re.search(r"故选\s*[：:]?\s*([A-D])(?![A-Za-z])", t)
            if m and qnum not in key:
                key[qnum] = (m.group(1), pg)
            if re.match(r"^\s*(?:【\s*答案\s*】|答案\s*[：:])", t):
                for m in re.finditer(r"(\d{1,2})\s*[．.、:：]\s*([A-D])(?![A-Za-z])", t):
                    key.setdefault(int(m.group(1)), (m.group(2), pg))
        if b["kind"] in ("p", "textbox"):
            if re.search(r"(\d{1,2}\s*[．.、:：]?\s*[A-D][\s\t|，,;；]*){3,}", t):
                for m in re.finditer(r"(\d{1,2})\s*[．.、:：]?\s*([A-D])(?![A-Za-z])", t):
                    key.setdefault(int(m.group(1)), (m.group(2), pg))
            for m in re.finditer(r"(\d{1,2})\s*[-—～~至]\s*(\d{1,2})\s*[:：．.、]?\s*([A-D]{2,})", t):
                a, z, letters = int(m.group(1)), int(m.group(2)), m.group(3)
                if z - a + 1 == len(letters):
                    for k, c in enumerate(letters):
                        key.setdefault(a + k, (c, pg))
            m = re.match(r"^\s*(\d{1,2})\s*[．.、]?\s*【\s*答案\s*】\s*([A-D])\b", t)
            if m:
                key.setdefault(int(m.group(1)), (m.group(2), pg))
        if b["kind"] == "table":
            rows = [[c.strip() for c in r.strip().strip("|").split("|")] for r in t.splitlines()
                    if r.strip() and not r.startswith("|---")]
            if "<table>" in t:
                rows = [[re.sub(r"<[^>]+>", "", c).strip() for c in re.findall(r"<td[^>]*>(.*?)</td>", r)]
                        for r in re.findall(r"<tr>(.*?)</tr>", t)]
            for i in range(len(rows) - 1):
                a, c = rows[i], rows[i + 1]
                nums = [x for x in a if x.isdigit()]
                lets = [x for x in c if re.fullmatch(r"[A-D]", x)]
                if len(nums) >= 3 and len(lets) >= 3:
                    aa = [x for x in a if x.isdigit() or "题号" in x]
                    cc = [x for x in c if re.fullmatch(r"[A-D]", x) or "答案" in x]
                    off_a = 1 if aa and not aa[0].isdigit() else 0
                    off_c = 1 if cc and not re.fullmatch(r"[A-D]", cc[0]) else 0
                    for x, y in zip(aa[off_a:], cc[off_c:]):
                        if x.isdigit() and re.fullmatch(r"[A-D]", y):
                            key.setdefault(int(x), (y, pg))
    return key


# ---------------------------------------------------------------- 渲染
def render_blocks(bs, residual_ids=None):
    residual_ids = residual_ids or {}
    out = []
    for b in bs:
        k = b["kind"]
        if k == "pagemark":
            out.append("<!-- 第%d页 %s -->" % (b["page"], b.get("class") or ""))
        elif k == "p":
            tag = "〔OCR候选〕" if b.get("ocr") else ("〔字体映射疑误〕" if b.get("garbled") else "")
            out.append(tag + b["text"])
        elif k == "textbox":
            rid = residual_ids.get(id(b))
            out.append("> 〔文本框%s〕" % ("|" + rid if rid else "") + b["text"].replace("\n", "\n> "))
        elif k == "table":
            out.append(b["text"])
        elif k == "image":
            rid = residual_ids.get(id(b), "")
            out.append("![图](%s)\n〔待转写:%s〕" % (b.get("path") or "", rid or "图"))
        elif k == "shape":
            rid = residual_ids.get(id(b), "")
            out.append("〔形状/框线/%s对象，待看图:%s〕" % (b.get("what", ""), rid))
        elif k == "scanpage":
            rid = residual_ids.get(id(b), "")
            out.append("![原页](%s)\n〔本页页型 %s：文字如有均为 OCR 候选或字体可疑，须模型对照原页出补丁:%s〕"
                       % (b.get("path") or "", b.get("class"), rid))
        elif k == "notes":
            out.append("〔演讲者备注〕" + b["text"])
    return "\n\n".join(x for x in out if x)


def page_span(pages):
    ps = sorted(p for p in pages if p)
    if not ps:
        return "?"
    return str(ps[0]) if ps[0] == ps[-1] else "%d—%d" % (ps[0], ps[-1])


# ---------------------------------------------------------------- 残留
def build_residual(blocks, pages, sid, role, rel, exam_id=None, start=0):
    """把图、扫描页、OCR 文字层页、乱码页、文本框、复杂表格、被剔除的不可见文字等列为 residual；
    返回 (rows, {id(block): rid})。文字字段超长时截断并标注，全文位置写在 full_text_ref。"""
    rows, ids = [], {}
    k = start
    ptype = {p["page"]: p for p in pages}

    def ref(pg):
        return "%s/source.full.md 第%s页" % (sid, pg)

    def ctx(i, step):
        j = i + step
        while 0 <= j < len(blocks):
            t = plain(blocks[j].get("text", "")).strip()
            if t:
                pg = blocks[j].get("page")
                # 前文取上一块末尾、后文取下一块开头；截断处标注
                if step < 0:
                    return ("〔前略〕" + t[-60:]) if len(t) > 60 else t
                return clip(t, 60, ref(pg))
            j += step
        return ""
    for i, b in enumerate(blocks):
        kind = None
        if b["kind"] == "image":
            kind = "picture_slide" if b.get("area", 0) >= 0.5 else (
                "vector_frame" if b.get("what") == "vector" else ("docx_image" if rel.lower().endswith(".docx") else "figure"))
            task = "transcribe_figure"
        elif b["kind"] == "scanpage":
            kind = "ocr_layer_page" if b.get("class") in OCR_LAYER_CLASSES else "scan_page"
            task = "proofread_patch"
        elif b["kind"] == "shape":
            kind, task = "vector_frame", "transcribe_figure"
        elif b["kind"] == "textbox":
            kind, task = "floating_textbox", "check_order"
        elif b["kind"] == "table" and (b.get("merged") or len(set(b.get("cols") or [0])) > 1):
            kind, task = "table_complex", "check_table"
        if not kind:
            continue
        k += 1
        rid = "R-%s-%04d" % (sid, k)
        ids[id(b)] = rid
        meta = ptype.get(b.get("page"), {})
        full = plain(b.get("text", ""))
        rows.append({"id": rid, "exam_id": exam_id, "qid": None, "source_id": sid, "role": role, "rel_path": rel,
                     "page": b.get("page"), "page_type": meta.get("page_type"), "fine_class": meta.get("fine_class"),
                     "kind": kind, "bbox": b.get("bbox"), "crop_path": b.get("path"),
                     "context_before": ctx(i, -1), "context_after": ctx(i, 1),
                     "draft_text": clip(full, 200, ref(b.get("page"))) or None,
                     "truncated": len(full) > 200, "full_text_ref": ref(b.get("page")) if len(full) > 200 else None,
                     "task": task, "affects_answering": None})
    for p in pages:
        if p.get("page_type") == "garbled_font":
            k += 1
            rows.append({"id": "R-%s-%04d" % (sid, k), "exam_id": exam_id, "qid": None, "source_id": sid, "role": role,
                         "rel_path": rel, "page": p["page"], "page_type": "garbled_font", "fine_class": p["fine_class"],
                         "kind": "garbled_symbols", "bbox": None, "crop_path": None, "context_before": "",
                         "context_after": "", "draft_text": None, "task": "proofread_patch",
                         "affects_answering": None})
    # 被剔除的文字（白底白字、低透明度水印、印厂排版标签）：同一原因同一文字合并一行，附全部页码与原文
    grp = collections.OrderedDict()
    for p in pages:
        for d_ in p.get("dropped") or []:
            key = (d_["reason"], d_["text"].strip())
            g = grp.setdefault(key, {"pages": [], "bbox": d_.get("bbox"), "ratio": d_.get("bg_nonwhite_ratio")})
            if p["page"] not in g["pages"]:
                g["pages"].append(p["page"])
    for (reason, text), g in grp.items():
        k += 1
        rows.append({"id": "R-%s-%04d" % (sid, k), "exam_id": exam_id, "qid": None, "source_id": sid, "role": role,
                     "rel_path": rel, "page": g["pages"][0], "pages": g["pages"],
                     "page_type": ptype.get(g["pages"][0], {}).get("page_type"),
                     "fine_class": ptype.get(g["pages"][0], {}).get("fine_class"),
                     "kind": "dropped_invisible_text", "reason": reason, "bbox": g["bbox"],
                     "bg_nonwhite_ratio": g["ratio"], "crop_path": None, "context_before": "", "context_after": "",
                     "draft_text": text, "truncated": False, "task": "adjudicate_dropped_text",
                     "affects_answering": None})
    return rows, ids


# ---------------------------------------------------------------- 自检工具
def options_check(text):
    pt = canon(plain(text))
    letters = re.findall(r"(?m)(?:^|[\s\t\u3000])(?:〔OCR候选〕)?([A-D])\s*[．.、]", pt)
    c = collections.Counter(letters)
    circ_ref = set(re.findall(r"(?m)^[A-D]\s*[．.、]\s*([①②③④⑤⑥]+)", pt))
    need = set("".join(circ_ref))
    body = "\n".join(l for l in pt.splitlines() if not re.match(r"^\s*(?:〔OCR候选〕)?[A-D]\s*[．.、]", l))
    have = set(re.findall(r"[①②③④⑤⑥]", body))
    return {"letters": "".join(sorted(c)), "complete": all(c.get(x, 0) >= 1 for x in "ABCD"),
            "duplicate": sorted(x for x, v in c.items() if v > 1),
            "circled_needed": "".join(sorted(need)), "circled_missing": "".join(sorted(need - have))}


def score_of(stem_text, n):
    pt = plain(stem_text)
    first = pt.strip().splitlines()[0] if pt.strip() else ""
    m = re.match(r"^\s*(?:〔OCR候选〕)?\d{1,2}\s*[．.、：:，,]?\s*[（(]\s*(\d{1,2})\s*分\s*[）)]", first)
    if m:
        return int(m.group(1)), "题首分值"
    subs = [int(x) for x in RE_SCORE.findall(pt)]
    if subs:
        return sum(subs), "小问/题末分值合计(%s)" % "+".join(map(str, subs))
    return None, None


RE_HEAD_SCORE = re.compile(r"^\s*(?:〔OCR候选〕)?(?:第\s*)?\d{1,2}\s*(?:题)?\s*[．.、]?\s*[（(]\s*(?:共\s*)?(\d{1,2})\s*分\s*[）)]")
RE_HEAD_TAIL = re.compile(r"^\s*(?:第\s*)?\d{1,2}\s*(?:题)?\s*[．.、](?!\s*(?:\d\s*[．.、]|[（(]\s*\d\s*[）)])).{8,}?[（(]\s*(\d{1,2})\s*分\s*[）)]\s*$")
RE_TOTAL_SCORE = re.compile(r"[（(]\s*共\s*(\d{1,2})\s*分\s*[）)]|本题\s*(?:共|满分)?\s*(\d{1,2})\s*分|满分\s*(\d{1,2})\s*分")


def heading_score(bs):
    """评分材料/答案区题首的本题总分：只认“16．（8分）”“（共8分）”“本题8分”，不累加细则里的逐点分。"""
    for b in bs:
        t = plain(b.get("text", "")).strip()
        m = RE_HEAD_SCORE.match(t)
        if m:
            return int(m.group(1))
        first = t.splitlines()[0] if t else ""
        m = RE_HEAD_TAIL.match(first)
        if m and len(RE_SCORE.findall(first)) == 1:
            return int(m.group(1))
        m = RE_TOTAL_SCORE.search(t[:80])
        if m:
            return int(next(g for g in m.groups() if g))
    return None


def table_check(bs):
    bad = []
    for b in bs:
        if b["kind"] == "table":
            cols = b.get("cols") or []
            if cols and len(set(cols)) > 1:
                bad.append({"page": b.get("page"), "cols": cols})
    return bad


def asset_check(md_text, base_dir):
    miss = []
    for m in re.finditer(r"!\[[^\]]*\]\(([^)]+)\)", md_text):
        p = m.group(1)
        if p and not os.path.isfile(os.path.join(base_dir, p)):
            miss.append(p)
    return miss


# ---------------------------------------------------------------- 单件转换
def _qmd(qid, exam_id, n, meta, sections, guide, check):
    L = ["# %s" % qid, "", "## 身份信息",
         "- question_id: `%s`" % qid, "- exam_id: `%s`" % (exam_id or "未指定"), "- 题号: %d" % n,
         "- 题型: %s" % meta["qtype"], "- 分值: %s" % meta["score_text"],
         "- 小问: %s" % (meta.get("subqs") or "无（或未标注）"),
         "- 来源页: %s" % meta["pages_text"],
         "- 生成: %s 确定性抽取；证据等级 `machine_extracted`（未经模型或人工核对，不得标为 verified）" % VERSION,
         "", "## 读取指引（机器可读）", ""]
    L += guide
    L.append("")
    # 同名小节合并成一个“## 标题”下的多个“### 来源”（bank.py 按标题建字典，重复标题会丢掉前面的节）
    merged = collections.OrderedDict()
    for title, sub in sections:
        merged.setdefault(title, []).extend(sub)
    for title, sub in merged.items():
        L.append("## %s" % title)
        L.append("")
        for head, body in sub:
            if head:
                L.append("### %s" % head)
                L.append("")
            L.append(body)
            L.append("")
    L += ["## 学生示例、教师批注与实得分", "",
          "- doc2md 不判定学生层（学生样卷多为图片，见 residual）；不得据此记 N/A。", "",
          "## 机器自检", "", "```json", json.dumps(check, ensure_ascii=False), "```",
          "", "> 机器自检通过只说明机械层没有发现问题，不等于内容已核验。", ""]
    return "\n".join(L)


def _filename_qhint(path):
    b = os.path.basename(path)
    m = re.search(r"(?:第)?(\d{1,2})\s*(?:题|[\.．]\s*(?:docx?|pptx?|pdf)$)", b)
    return int(m.group(1)) if m else None


def convert_file(path, role="原卷", exam_id=None, out_dir=DEFAULT_OUT, name=None, fmt=None, sid=None, rel_path=None,
                 ocr_dirs=None, run_ocr_pages=False, render_dpi=110, known_numbers=None, stems=None, subjective=None,
                 write=True):
    """单个原件 → 候选目录。返回结果 dict（含 blocks、split、residual、selfcheck，供整卷模式合并）。"""
    if write and not out_dir:
        raise SystemExit("没有输出目录：请给 --out-dir 或设置环境变量 DOC2MD_OUT")
    t0 = time.time()
    fmt = fmt or detect_format(path)
    digest = sha256(path)
    sid = sid or sid_of(path, digest)
    rel = rel_path or path
    name = name or ("%s_%s" % (exam_id, sid) if exam_id else sid)
    odir = os.path.join(out_dir or "", name)
    if write:
        os.makedirs(os.path.join(odir, "questions"), exist_ok=True)
    blocks, flags, pages, info = extract_source(
        path, fmt, sid, asset_root=odir if write else None, ocr_dirs=ocr_dirs, render_dpi=render_dpi,
        ocr_run_dir=os.path.join(odir, "ocr_run") if run_ocr_pages else None, work_dir=os.path.join(odir, "legacy"))
    info.update({"role": role, "rel_path": rel, "exam_id": exam_id})
    layer = ROLE_LAYER.get(role, "aux")
    residual, rids = build_residual(blocks, pages, sid, role, rel, exam_id)
    res = {"info": info, "blocks": blocks, "pages": pages, "flags": dict(flags), "layer": layer,
           "residual": residual, "rids": rids, "odir": odir, "questions": {},
           "known_numbers": sorted(known_numbers) if known_numbers else None}
    if layer in ("paper", "teacher"):
        sp = split_paper(blocks)
        res["split"] = sp
        res["key"] = answer_key(sp["answer_region"]) if sp["answer_region"] else {}
        for n, q in sp["questions"].items():
            k2 = answer_key(q["answer"], qnum=n)
            if n in k2:
                res["key"].setdefault(n, k2[n])
        subj = [n for n in sp["questions"] if not _is_choice(n, sp)]
        own = stem_grams(sp)
        pre, aq, ev = split_rubric(sp["answer_region"], known=set(sp["questions"]) or None, stems=own,
                                   answer_region=True) \
            if sp["answer_region"] else ([], {}, [])
        for n, bs in aq.items():
            for kk, vv in answer_key(bs, qnum=n).items():
                res["key"].setdefault(kk, vv)
        res["embedded_answers"] = aq
        res["split"]["events"] += ev
    else:
        pre, qs, ev = split_rubric(blocks, known=known_numbers, hint=_filename_qhint(path), stems=stems,
                                   subjective=subjective if fmt == "pptx" else None,
                                   choice=(set(known_numbers) - set(subjective)) if known_numbers and subjective else None)
        for e_ in ev:
            if e_.get("event") == "rubric_sections_by_restart" and e_.get("inferred"):
                res["split_uncertain"] = sorted(set(res.get("split_uncertain", [])) | set(e_["inferred"]))
        # 首个显式题号之前的大段正文：若前一号是原卷主观题且本件没有它，按顺序归给它（标 split_uncertain）
        if subjective and qs:
            n0 = next(iter(qs))
            body = sum(len(plain(b.get("text", ""))) for b in pre if b["kind"] in ("p", "table", "textbox"))
            prev_q = max([x for x in subjective if x < n0] or [0])
            if prev_q and prev_q not in qs and body > 200:
                cut = next((i for i, b in enumerate(pre) if b["kind"] == "pagemark" and i > 0 and b.get("page", 0) > 1), 0)
                cut = 0 if cut == 0 else cut
                keep_front, moved = pre[:cut], pre[cut:]
                if not moved:
                    keep_front, moved = [], pre
                new = collections.OrderedDict([(prev_q, moved)])
                new.update(qs)
                qs, pre = new, keep_front
                ev.append({"event": "rubric_pre_assigned_by_order", "q": prev_q, "chars": body})
                res["split_uncertain"] = res.get("split_uncertain", []) + [prev_q]
        res["split"] = {"front": pre, "questions": qs, "events": ev}
        res["key"] = answer_key(blocks)
        for n, bs in qs.items():
            for kk, vv in answer_key(bs, qnum=n).items():
                res["key"].setdefault(kk, vv)
    # 页映射：每页至少归到一题或卷首
    mapped = collections.defaultdict(set)
    sq = res["split"]["questions"]
    for n, q in sq.items():
        bl = (q["stem"] + q["answer"]) if isinstance(q, dict) else q
        for b in bl:
            if b.get("page") and b["kind"] != "pagemark":
                mapped[b["page"]].add("Q%d" % n)
    for b in res["split"].get("front", []) + res["split"].get("answer_region", []):
        if b.get("page") and b["kind"] != "pagemark":
            mapped[b["page"]].add("front" if b in res["split"].get("front", []) else "answer_region")
    for p in pages:
        p["mapped_to"] = sorted(mapped.get(p["page"], set()))
    for r in residual:
        qs_ = mapped.get(r["page"], set())
        r["qid_candidates"] = sorted(qs_)
    res["selfcheck"] = selfcheck_single(res)
    info["seconds_total"] = round(time.time() - t0, 3)
    if write:
        write_single(res, exam_id, name)
    return res


def stem_grams(sp):
    return {n: _grams(canon(norm_display(plain(render_blocks(q["stem"]))))) for n, q in sp["questions"].items()}


def _section_of(n, sp):
    secs = [x for x in sp.get("sections", []) if x.get("start_q") and x["start_q"] <= n and
            (x.get("count") or x.get("total") or x.get("per_q"))]
    return secs[-1] if secs else None


def _is_choice(n, sp):
    sec = _section_of(n, sp)
    if sec is not None:
        if sec.get("count") and n >= sec["start_q"] + sec["count"]:
            sec = None
        else:
            return bool(sec.get("per_q") or sec.get("choice"))
    q = sp["questions"].get(n)
    if q:
        return options_check(render_blocks(q["stem"]))["complete"]
    return False


def _per_q(n, sp):
    sec = _section_of(n, sp)
    if sec and sec.get("per_q") and n < sec["start_q"] + (sec.get("count") or 99):
        return sec["per_q"]
    return None


def declared_count(sections):
    """卷首各部分“本部分共N题…共M分”：只有各部分分值合计正好 100（即各部分都声明了题数）时，返回题数合计。"""
    secs = [x for x in sections if x.get("count") and x.get("total")]
    if secs and sum(x["total"] for x in secs) == 100:
        return sum(x["count"] for x in secs)
    return None


def _answer_text(bs):
    return plain(render_blocks([b for b in bs if b["kind"] != "pagemark"]))


RE_LATER_HEAD = [RE_R_STRONG[0], RE_R_STRONG[1], RE_R_STRONG[2], RE_R_STRONG[7]]   # 第N题 / N题 / 【N题】/ N.评分细则
RE_LATER_SCORE_HEAD = RE_R_STRONG[3]                                  # N．（M分）：常是细则里引用的“相关试题”


def block_mixing_issues(n, bs, label, known=None, soft=None):
    """一块答案里出现别题的“N．【分析】”题首、两个以上不同题号题首或多个“故选”：切题串块（返回硬失败列表）。
    known 给定时（评分材料/参考答案），块里还有后面题号的“第18题/18题/【18题】”也算串块；
    只有“18．（8分）”式题首的记到 soft（警告），因为细则常引用别卷的“相关试题”。"""
    txt = _answer_text(bs)
    heads = {int(x) for x in RE_ANS_HEAD_N_ML.findall(txt)}
    out = []
    if known:
        later, later_soft = set(), set()
        for line in txt.splitlines():
            t_ = line.strip()
            for rx in RE_LATER_HEAD:
                m = rx.match(t_)
                if m and int(m.group(1)) > n and int(m.group(1)) in known:
                    later.add(int(m.group(1)))
            m = RE_LATER_SCORE_HEAD.match(t_)
            if m and int(m.group(1)) > n and int(m.group(1)) in known:
                later_soft.add(int(m.group(1)))
        if later:
            out.append("%sQ%d 评分块里还有后面题号 %s 的题首（切题串块）" % (label, n, sorted(later)))
        if later_soft - later and soft is not None:
            soft.append("%sQ%d 评分块里有“%s．（M分）”式题首，可能是串块，也可能是引用的相关试题，须人工看" % (
                label, n, "/".join(map(str, sorted(later_soft - later)))))
    if len(heads) >= 2 or (heads and n not in heads):
        out.append("%sQ%d 答案块含题号 %s 的答案题首（切题串块）" % (label, n, sorted(heads)))
    k = len(re.findall(r"故选", txt))
    if k >= 2:
        out.append("%sQ%d 答案块出现 %d 处“故选”（疑似多题答案并入一块）" % (label, n, k))
    return out


def answer_block_issues(res):
    """原卷/教师版内附答案：串块，以及“同类题有答案块、本题却没有”。返回硬失败列表。"""
    sp = res["split"]
    qs = sp["questions"]
    emb = res.get("embedded_answers", {})
    hard, have = [], {}
    for n, q in qs.items():
        bs = list(q.get("answer", [])) + list(emb.get(n, []))
        have[n] = any(b["kind"] != "pagemark" and plain(b.get("text", "")).strip() for b in bs)
        if have[n]:
            hard += block_mixing_issues(n, bs, "")
    for kind in (True, False):
        group = [n for n in qs if _is_choice(n, sp) == kind]
        if not any(have[n] for n in group):
            continue
        for n in group:
            if not have[n] and not (kind and n in res.get("key", {})):
                hard.append("Q%d 本件有同类%s题的答案块，但本题既无答案块也无答案键" % (n, "选择" if kind else "非选择"))
    return hard


def selfcheck_single(res):
    sc = {"source": {k: res["info"].get(k) for k in ("source_id", "rel_path", "format", "role", "sha256",
                                                    "page_types", "fine_classes")},
          "hard_fail": [], "warn": [], "note": "机器自检通过只说明机械层没有发现问题，不等于内容已核验。"}
    sq = res["split"]["questions"]
    nums = list(sq.keys())
    if res["layer"] in ("paper", "teacher"):
        contiguous = nums == list(range(1, len(nums) + 1))
        sc["questions"] = len(nums)
        sc["numbers_contiguous"] = contiguous
        if not contiguous:
            sc["hard_fail"].append("题号不连续：%s" % nums)
        sp = res["split"]
        dec_n = declared_count(sp["sections"])
        sc["declared_question_count"] = dec_n
        if not nums:
            sc["hard_fail"].append("未切出任何题")
        if not dec_n:
            sc["warn"].append("卷首未声明完整的“本部分共N题/共M分”，无法由本件推算应有题数")
        total, unresolved, detail = 0, [], {}
        for n, q in sq.items():
            stem = render_blocks(q["stem"])
            if _is_choice(n, sp):
                oc = options_check(stem)
                if not oc["complete"]:
                    sc["hard_fail"].append("Q%d 选项不齐：%s" % (n, oc["letters"]))
                if oc["duplicate"]:
                    sc["hard_fail"].append("Q%d 选项字母重复：%s（疑似两题选项并入一块）" % (n, "".join(oc["duplicate"])))
                if oc["circled_missing"]:
                    sc["warn"].append("Q%d 选项引用的题肢缺失：%s" % (n, oc["circled_missing"]))
                per = _per_q(n, sp)
                detail[n] = per
                if per is None:
                    unresolved.append(n)
                else:
                    total += per
                if n not in res.get("key", {}):
                    sc["warn"].append("Q%d 本件未见答案键" % n) if res["layer"] == "teacher" else None
            else:
                s_, basis = score_of(stem, n)
                detail[n] = s_
                if s_ is None:
                    unresolved.append(n)
                else:
                    total += s_
        sc["score_by_question"] = detail
        sc["score_total_resolved"] = total
        sc["score_unresolved"] = unresolved
        declared = [s for s in sp["sections"] if s.get("total")]
        sc["sections_declared"] = [{"start_q": s["start_q"], "count": s["count"], "per_q": s["per_q"],
                                    "total": s["total"]} for s in declared]
        dec_total = sum(s["total"] for s in declared) if declared else None
        sc["score_declared_total"] = dec_total
        if not unresolved and nums and total != 100:
            sc["hard_fail"].append("分值合计 %d ≠ 100" % total)
        if dec_n and dec_n != len(nums):
            msg = "题数 %d 与卷首“本部分共N题”合计 %d 不符（末题号 %s）" % (len(nums), dec_n, nums[-1] if nums else "无")
            if not unresolved and total == 100:
                sc["warn"].append(msg + "；但逐题分值齐全且合计 100，疑原件卷首题数印误，须人工看原页")
            else:
                sc["hard_fail"].append(msg)
        if unresolved:
            if dec_total == 100 and total < 100:
                sc["warn"].append("Q%s 分值本件未印，按卷首分部总分 %s=100 核对（整卷模式再用 E1/E3 补）"
                                  % (",".join(map(str, unresolved)), "+".join(str(s["total"]) for s in declared)))
            else:
                sc["hard_fail"].append("Q%s 分值未解析，且卷首分部总分为 %s（不是 100），总分无法核对" % (
                    ",".join(map(str, unresolved)), dec_total))
        sc["answer_block_issues"] = answer_block_issues(res)
        sc["hard_fail"] += sc["answer_block_issues"]
        sc["answer_region"] = bool(res["split"].get("answer_region"))
        sc["answer_keys"] = {str(k): v[0] for k, v in sorted(res.get("key", {}).items())}
    else:
        sc["rubric_blocks"] = nums
        sc["answer_keys"] = {str(k): v[0] for k, v in sorted(res.get("key", {}).items())}
        if not nums and not res.get("key"):
            sc["warn"].append("未切出任何题号块，也没有答案键（split_uncertain）")
        mix = []
        kn = set(res.get("known_numbers") or sq.keys())
        for n, bs in sq.items():
            mix += block_mixing_issues(n, bs, "%s " % (res["info"].get("role") or ""), known=kn, soft=sc["warn"])
        sc["answer_block_issues"] = mix
        # 讲评/角色未定材料常在一页里讨论多道题，串块只记警告；评分材料与参考答案记硬失败
        (sc["hard_fail"] if res["layer"] in ("E1", "E3") else sc["warn"]).extend(mix)
    if res.get("split_uncertain"):
        sc["warn"].append("Q%s 的评分块没有显式题号，按顺序归入（split_uncertain，须人工确认）"
                          % ",".join(map(str, res["split_uncertain"])))
        res["residual"].append({"id": "R-%s-split" % res["info"]["source_id"], "exam_id": res["info"].get("exam_id"),
                                "qid": None, "source_id": res["info"]["source_id"], "role": res["info"]["role"],
                                "rel_path": res["info"]["rel_path"], "page": None, "kind": "split_uncertain",
                                "task": "adjudicate", "draft_text": "无显式题号的评分块归入 Q%s" % res["split_uncertain"],
                                "qid_candidates": ["Q%d" % q for q in res["split_uncertain"]], "affects_answering": None})
    unmapped = [p["page"] for p in res["pages"] if not p.get("mapped_to") and p.get("page_type") != "empty"]
    sc["pages_unmapped"] = unmapped
    if unmapped:
        sc["warn"].append("来源页未映射到任何题或卷首：%s" % unmapped)
    tb = table_check(res["blocks"])
    if tb:
        sc["warn"].append("表格行列不齐 %d 处（已列 residual table_complex）" % len(tb))
    sc["residual_items"] = len(res["residual"])
    sc["residual_by_kind"] = dict(collections.Counter(r["kind"] for r in res["residual"]))
    return sc


def _section_head(info, layer_tag, pages, classes):
    return "来源：`%s`（角色 %s / %s，第%s页，页型 %s）" % (
        info.get("rel_path"), info.get("role"), layer_tag, page_span(pages),
        "、".join("%s×%d" % (k, v) for k, v in sorted(collections.Counter(classes).items())))


def _split_direct(bs):
    """把一题的块按页型分成 native 直出与 OCR/乱码候选两组（保持原顺序）。"""
    direct, ocr = [], []
    for b in bs:
        if b.get("ocr") or b["kind"] == "scanpage" and b.get("class") not in DIRECT_CLASSES:
            ocr.append(b)
        else:
            direct.append(b)
    return direct, ocr


def question_sections(res, n, part):
    """返回某题在本件中的小节列表 [(title, [(head, body)])]。part: stem|answer|rubric|embedded"""
    info, layer = res["info"], res["layer"]
    if part in ("stem", "answer"):
        bs = res["split"]["questions"][n][part]
    elif part == "embedded":
        bs = res.get("embedded_answers", {}).get(n, [])
    else:
        bs = res["split"]["questions"].get(n, [])
    bs = [b for b in bs if b["kind"] != "pagemark" or True]
    if not [b for b in bs if b["kind"] != "pagemark"]:
        return []
    out = []
    direct, ocr = _split_direct(bs)
    for group, is_ocr in ((direct, False), (ocr, True)):
        body_bs = [b for b in group if b["kind"] != "pagemark"]
        if not body_bs:
            continue
        pages = sorted({b["page"] for b in body_bs if b.get("page")})
        classes = [b.get("class") or "?" for b in body_bs if b["kind"] != "pagemark"]
        marks = []
        seen = set()
        for b in group:
            if b.get("page") and b["page"] not in seen and b["kind"] != "pagemark":
                seen.add(b["page"])
        body = render_blocks(_with_pagemarks(group), res["rids"])
        if part == "stem":
            title = LAYER_TITLE["ocr"] if is_ocr else LAYER_TITLE["teacher" if layer == "teacher" else "paper"]
        elif part in ("answer", "embedded"):
            title = LAYER_TITLE["E3"]
        else:
            title = LAYER_TITLE.get(layer, LAYER_TITLE["aux"])
        tag = "ocr_candidate" if is_ocr else "native"
        if part == "answer":
            tag += "，教师版内附答案/解析（E3）"
        if part == "embedded":
            tag += "，卷内答案区（E3，非独立答案源）"
        if part == "rubric":
            tag += "，配对方式：题号起首（machine_split，未核）"
        out.append((title, [(_section_head(info, tag, pages, classes), body)]))
    return out


def _with_pagemarks(bs):
    out, last = [], None
    for b in bs:
        if b["kind"] == "pagemark":
            continue
        if b.get("page") != last:
            out.append({"kind": "pagemark", "page": b.get("page") or 0, "class": b.get("class")})
            last = b.get("page")
        out.append(b)
    return out


def question_meta(res, n, extra_score=None):
    sp = res["split"]
    q = sp["questions"][n]
    stem = render_blocks(q["stem"])
    choice = _is_choice(n, sp)
    if choice:
        per = _per_q(n, sp)
        score, basis = per, "卷首“每题%s分”" % per if per else None
    else:
        score, basis = score_of(stem, n)
    if score is None and extra_score:
        score, basis = extra_score
    subqs = re.findall(r"(?m)^\s*(?:〔OCR候选〕)?[（(]\s*(\d)\s*[）)]", plain(stem))
    classes = collections.Counter(b.get("class") or "?" for b in q["stem"] if b["kind"] != "pagemark")
    return {"qtype": "选择题" if choice else "非选择题", "score": score,
            "score_text": ("（%d分）（依据：%s）" % (score, basis)) if score is not None else "〔未解析〕",
            "subqs": "".join("（%s）" % s for s in dict.fromkeys(subqs)) or None,
            "pages_text": "%s（页型 %s）" % (page_span(q["pages"]), "、".join("%s×%d" % kv for kv in classes.items())),
            "choice": choice, "stem_text": stem}


def write_single(res, exam_id, name):
    odir = res["odir"]
    info = res["info"]
    with open(os.path.join(odir, "source.full.md"), "w", encoding="utf-8") as fh:
        fh.write("# %s\n\n> source_id: %s；角色：%s；sha256：%s；%s；证据等级 machine_extracted\n\n" % (
            os.path.basename(info["path"]), info["source_id"], info["role"], info["sha256"], VERSION))
        fh.write(render_blocks(res["blocks"], res["rids"]))
    with open(os.path.join(odir, "pages.jsonl"), "w", encoding="utf-8") as fh:
        for p in res["pages"]:
            fh.write(json.dumps(dict(p, source_id=info["source_id"]), ensure_ascii=False) + "\n")
    with open(os.path.join(odir, "residual.jsonl"), "w", encoding="utf-8") as fh:
        for r in res["residual"]:
            fh.write(json.dumps(r, ensure_ascii=False) + "\n")
    base = exam_id or info["source_id"]
    sq = res["split"]["questions"]
    for n in sq:
        qid = "%s-Q%d" % (base, n)
        if res["layer"] in ("paper", "teacher"):
            meta = question_meta(res, n)
            secs = question_sections(res, n, "stem") + question_sections(res, n, "answer") + \
                question_sections(res, n, "embedded")
            if n in res.get("key", {}):
                v, pg = res["key"][n]
                secs.append((LAYER_TITLE["E3"], [("来源：`%s`（答案键，第%s页）" % (info["rel_path"], pg),
                                                  "- 选择题答案键：第%d题 **%s**" % (n, v))]))
        else:
            meta = {"qtype": "（本件为%s，不判题型）" % info["role"], "score_text": "〔见原卷〕", "subqs": None,
                    "pages_text": page_span({b.get("page") for b in sq[n] if b.get("page")}), "choice": False}
            secs = question_sections(res, n, "rubric")
            if n in res.get("key", {}):
                v, pg = res["key"][n]
                secs.append((LAYER_TITLE["E1"] if res["layer"] == "E1" else LAYER_TITLE["E3"],
                             [("来源：`%s`（答案键，第%s页）" % (info["rel_path"], pg), "- 选择题答案键：第%d题 **%s**" % (n, v))]))
        guide = _guide(secs, res)
        chk = {"source_id": info["source_id"], "page_types": info["page_types"],
               "residual": [r["id"] for r in res["residual"] if ("Q%d" % n) in r.get("qid_candidates", [])]}
        md = _qmd(qid, exam_id, n, meta, secs, guide, chk)
        with open(os.path.join(odir, "questions", qid + ".md"), "w", encoding="utf-8") as fh:
            fh.write(md)
    sc = res["selfcheck"]
    miss = []
    for f in os.listdir(os.path.join(odir, "questions")):
        miss += asset_check(open(os.path.join(odir, "questions", f), encoding="utf-8").read(), odir + "/questions/..")
    sc["missing_asset_refs"] = sorted(set(miss))
    if miss:
        sc["hard_fail"].append("图片引用失效 %d 处" % len(set(miss)))
    json.dump(dict(sc, info=info), open(os.path.join(odir, "selfcheck.json"), "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)


def _guide(secs, res, note=None):
    titles = [t for t, _ in secs]
    stem = next((t for t in titles if t in (LAYER_TITLE["paper"], LAYER_TITLE["teacher"])), None)
    g = []
    if stem:
        g.append("- 采用题面小节: `%s`" % stem)
        g.append("- 采用题面来源性质: %s / native（machine_extracted，未核）" % res["info"]["role"])
    elif LAYER_TITLE["ocr"] in titles:
        g.append("- 采用题面小节: 无（本题题面只有 OCR 候选；须模型对照原页出补丁后改用 `题目原文（原卷·视觉核校转写）`）")
    if LAYER_TITLE["ocr"] in titles:
        g.append("- 备选题面小节: `%s`" % LAYER_TITLE["ocr"])
    g.append("- 不得当作题面: `%s`、`%s`、`%s`" % (LAYER_TITLE["E1"], LAYER_TITLE["E3"], LAYER_TITLE["lecture"]))
    if note:
        g.append("- 题面来源风险: %s" % note)
    g.append("- 残留（须模型看图，见 residual.jsonl）: 本题 machine 自检 JSON 的 residual 字段")
    return g


# ---------------------------------------------------------------- 整卷
def load_manifest():
    return json.load(open(MANIFEST, encoding="utf-8"))


def locate_source(rel):
    """主清单 rel_path → 本机路径：原料根、桌面，再到题库根（题库内的“可读派生见证”等）。"""
    for base in (RAW, DESKTOP, BANK):
        p = os.path.join(base, rel)
        if os.path.isfile(p):
            return p
    return None


def exam_sources(exam_id, man=None):
    man = man or load_manifest()
    task = next((t for t in man["tasks"] if t["exam_id"] == exam_id), None)
    if not task:
        raise SystemExit("转换主清单里没有 %s" % exam_id)
    out = []
    for role, lst in task["sources"].items():
        for s in lst:
            out.append(dict(s, role_key=role, path=locate_source(s["rel_path"])))
    return task, out


PAPER_ROLE_KEYS = ("原卷", "教师版", "混合载体", "教师版/转录题源")


def bank_question_count(exam_id):
    """题库该卷从 Q1 起连续的题文件数（历史兼容文件如 Q0/Q24/Q25 不算）；没有则 None。"""
    d = os.path.join(BANK, "questions", exam_id)
    if not os.path.isdir(d):
        return None
    nums = set()
    for f in os.listdir(d):
        m = re.match(r"^%s-Q(\d+)\.md$" % re.escape(exam_id), f)
        if m:
            nums.add(int(m.group(1)))
    k = 0
    while k + 1 in nums:
        k += 1
    return k or None


def expected_counts(task, exam_id):
    """应有题数的几路来源：主清单 question_index_count / question_md_count、题库题文件数。"""
    out = collections.OrderedDict()
    for k in ("question_index_count", "question_md_count"):
        v = task.get(k)
        if isinstance(v, int) and v > 0:
            out[k] = v
    b = bank_question_count(exam_id)
    if b:
        out["bank_question_files"] = b
    return out


def _paper_rank(r, expect):
    """原卷/教师版里选题面来源：先看切题是否完整（题数与应有题数相符、题号连续），再看文字层占比，
    最后才看角色（原卷优先于教师版）。对应题库 README 的“原卷 native > 教师版 native > 原卷 OCR 候选”。"""
    nums = list(r["split"]["questions"])
    n = len(nums)
    pts = collections.Counter(p["page_type"] for p in r["pages"])
    tot = sum(pts.values()) or 1
    share = (pts.get("text_reliable", 0) + pts.get("mixed", 0)) / float(tot)
    exp = set(expect.values()) | ({declared_count(r["split"]["sections"])} - {None})
    count_ok = (n in exp) if exp else n > 0
    bucket = 0 if share >= 0.8 else (1 if share >= 0.3 else 2)
    return (0 if count_ok else 1, 0 if nums == list(range(1, n + 1)) and n else 1, bucket,
            0 if r["layer"] == "paper" else 1, -share, -n)


def _convert_or_fail(s, exam_id, edir, run_ocr_pages, render_dpi, write, failed, residual, **kw):
    try:
        return convert_file(s["path"], s["role_key"], exam_id, os.path.join(edir, "sources"), name=s.get("source_id"),
                            fmt=s.get("format"), sid=s.get("source_id"), rel_path=s["rel_path"],
                            run_ocr_pages=run_ocr_pages, render_dpi=render_dpi, write=write, **kw)
    except Exception as e:
        # 单份原件打不开不拖垮整卷：记硬失败与 residual，其余来源照常转换
        failed.append({"rel_path": s["rel_path"], "role": s["role_key"], "error": "%s: %s" % (type(e).__name__, e)})
        residual.append({"id": "R-failed-%d" % len(failed), "exam_id": exam_id, "kind": "source_unreadable",
                         "rel_path": s["rel_path"], "role": s["role_key"], "task": "adjudicate",
                         "draft_text": "%s: %s" % (type(e).__name__, e)})
        return None


def convert_exam(exam_id, out_dir=DEFAULT_OUT, run_ocr_pages=False, render_dpi=110, write=True):
    if not out_dir:
        raise SystemExit("没有输出目录：请给 --out-dir 或设置环境变量 DOC2MD_OUT")
    t0 = time.time()
    task, srcs = exam_sources(exam_id)
    edir = os.path.join(out_dir, exam_id)
    results, residual, missing, failed = [], [], [], []
    for s in srcs:
        if not s["path"]:
            missing.append({"rel_path": s["rel_path"], "role": s["role_key"]})
            residual.append({"id": "R-missing-%d" % len(missing), "exam_id": exam_id, "kind": "missing_source",
                             "rel_path": s["rel_path"], "role": s["role_key"], "task": "adjudicate"})
            continue
        if s["role_key"] in PAPER_ROLE_KEYS:
            r = _convert_or_fail(s, exam_id, edir, run_ocr_pages, render_dpi, write, failed, residual)
            if r:
                results.append(r)
    expect = expected_counts(task, exam_id)
    ranked = sorted(results, key=lambda r: _paper_rank(r, expect))
    paper = ranked[0] if ranked else None
    cand_rows = [{"source_id": r["info"]["source_id"], "role": r["info"]["role"], "rel_path": r["info"]["rel_path"],
                  "questions": len(r["split"]["questions"]), "rank_key": list(_paper_rank(r, expect)),
                  "page_types": r["info"]["page_types"]} for r in ranked]
    known = set(paper["split"]["questions"]) if paper else None
    stems = stem_grams(paper["split"]) if paper else None
    subjective = {n for n in known if not _is_choice(n, paper["split"])} if paper else None
    for s in srcs:
        if not s["path"] or s["role_key"] in PAPER_ROLE_KEYS:
            continue
        r = _convert_or_fail(s, exam_id, edir, run_ocr_pages, render_dpi, write, failed, residual,
                             known_numbers=known, stems=stems, subjective=subjective)
        if r:
            results.append(r)
    for r in results:
        residual += r["residual"]
    pre = []
    if missing:
        pre.append("原件缺失 %d 份" % len(missing))
    if failed:
        pre.append("原件打不开 %d 份：%s" % (len(failed), "；".join(f["error"] for f in failed)))
    merged = merge_exam(exam_id, task, paper, results, edir, write, expect=expect, paper_candidates=cand_rows,
                        pre_hard_fail=pre)
    sc = merged["selfcheck"]
    sc["missing_sources"] = missing
    sc["failed_sources"] = failed
    sc["seconds"] = round(time.time() - t0, 2)
    for r in residual:
        if r.get("qid") is None and r.get("qid_candidates"):
            qn = [c for c in r["qid_candidates"] if c.startswith("Q")]
            r["qid"] = "%s-%s" % (exam_id, qn[0]) if len(qn) == 1 else None
    if write:
        os.makedirs(edir, exist_ok=True)
        with open(os.path.join(edir, "residual.jsonl"), "w", encoding="utf-8") as fh:
            for r in residual:
                fh.write(json.dumps(r, ensure_ascii=False) + "\n")
        with open(os.path.join(edir, "pages.jsonl"), "w", encoding="utf-8") as fh:
            for r in results:
                for p in r["pages"]:
                    fh.write(json.dumps(dict(p, source_id=r["info"]["source_id"], role=r["info"]["role"]),
                                        ensure_ascii=False) + "\n")
        json.dump(sc, open(os.path.join(edir, "selfcheck.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    merged["results"] = results
    merged["residual"] = residual
    return merged


def merge_exam(exam_id, task, paper, results, edir, write=True, expect=None, paper_candidates=None,
               pre_hard_fail=None):
    sc = {"exam_id": exam_id, "generator": VERSION, "hard_fail": list(pre_hard_fail or []), "warn": [],
          "note": "机器自检通过只说明机械层没有发现问题，不等于内容已核验；证据等级 machine_extracted。",
          "sources": [dict(r["info"], selfcheck_hard_fail=r["selfcheck"]["hard_fail"]) for r in results],
          "paper_candidates": paper_candidates or []}
    if not paper:
        sc["hard_fail"].append("没有可用的原卷/教师版来源，无法切题")
        return {"selfcheck": sc, "questions": {}}
    sq = paper["split"]["questions"]
    nums = list(sq)
    expect = expect if expect is not None else expected_counts(task, exam_id)
    sc["paper_source"] = paper["info"]["source_id"]
    sc["paper_role"] = paper["info"]["role"]
    sc["questions"] = len(nums)
    sc["expected_counts"] = dict(expect)
    sc["numbers_contiguous"] = bool(nums) and nums == list(range(1, len(nums) + 1))
    if not nums:
        sc["hard_fail"].append("题面来源未切出任何题")
    elif not sc["numbers_contiguous"]:
        sc["hard_fail"].append("题号不连续：%s" % nums)
    hit = [k for k, v in expect.items() if v == len(nums)]
    if expect and not hit:
        sc["hard_fail"].append("题数 %d 与应有题数都不符：%s" % (len(nums), dict(expect)))
    elif expect and len(hit) < len(expect):
        sc["warn"].append("题数 %d 与 %s 相符，但与 %s 不同（主清单计数可能过期）" % (
            len(nums), "、".join(hit), {k: v for k, v in expect.items() if k not in hit}))
    dec_n = declared_count(paper["split"]["sections"])
    sc["declared_question_count"] = dec_n
    guide_note = None
    if paper["layer"] != "paper":
        others = [r["info"]["rel_path"] for r in results if r is not paper and r["layer"] == "paper"]
        guide_note = "题面取自%s（教师版题面须回原卷核实）%s" % (
            paper["info"]["role"], "；原卷另有 %s（未合并，见 sources/）" % "、".join(others) if others else "；本卷无原卷来源")
        if others:
            sc["warn"].append("题面取自%s（切题更完整或文字层更可靠）；原卷 %s 未合并，题面须回原卷核实" % (
                paper["info"]["role"], "、".join(others)))
    # 题面来源自身的答案块串题/缺块，以及评分材料/参考答案的串块，都是整卷硬失败
    sc["answer_block_issues"] = list(paper["selfcheck"].get("answer_block_issues", []))
    for r in results:
        if r["layer"] in ("E1", "E3"):
            sc["answer_block_issues"] += ["%s：%s" % (r["info"]["source_id"], x)
                                          for x in r["selfcheck"].get("answer_block_issues", [])]
        if r is not paper:
            sc["warn"] += ["%s：%s" % (r["info"]["source_id"], w) for w in r["selfcheck"].get("warn", []) if "串块" in w]
    sc["hard_fail"] += sc["answer_block_issues"]
    keys, rub_score = {}, {}
    for r in results:
        for n, (v, pg) in r.get("key", {}).items():
            keys.setdefault(n, []).append((v, r["info"]["source_id"], r["info"]["role"], pg))
        if r["layer"] not in ("paper", "teacher"):
            for n, bs in r["split"]["questions"].items():
                v = heading_score([b for b in bs if b["kind"] != "pagemark"][:2])
                if v and n not in rub_score:
                    rub_score[n] = (v, "%s（%s）题首分值" % (r["info"]["role"], r["info"]["source_id"]))
        for n, bs in r.get("embedded_answers", {}).items():
            v = heading_score([b for b in bs if b["kind"] != "pagemark"][:2])
            if v and n not in rub_score:
                rub_score[n] = (v, "卷内答案区题首分值")
    total, unresolved, detail, qout = 0, [], {}, {}
    route_basis = {}
    for n in nums:
        meta = question_meta(paper, n, rub_score.get(n))
        detail[n] = meta["score"]
        if meta["score"] is None:
            unresolved.append(n)
        else:
            total += meta["score"]
        secs = question_sections(paper, n, "stem") + question_sections(paper, n, "answer") + \
            question_sections(paper, n, "embedded")
        for r in results:
            if r is paper:
                continue
            if r["layer"] in ("paper", "teacher"):
                continue  # 另一份原卷（如同卷 DOCX+PDF）不合并，单件目录里有
            if n in r["split"]["questions"]:
                secs += question_sections(r, n, "rubric")
        ks = keys.get(n, [])
        if ks:
            lines = ["- 选择题答案键：第%d题 **%s**（%s，第%s页）" % (n, v, role, pg) for v, s_, role, pg in ks]
            secs.append((LAYER_TITLE["E3"], [("答案键（machine_extracted）", "\n".join(lines))]))
        if meta["choice"]:
            oc = options_check(meta["stem_text"])
            if not oc["complete"]:
                sc["hard_fail"].append("Q%d 选项不齐：%s" % (n, oc["letters"] or "无"))
            if oc["duplicate"]:
                sc["hard_fail"].append("Q%d 选项字母重复：%s（疑似两题选项并入一块）" % (n, "".join(oc["duplicate"])))
            if oc["circled_missing"]:
                sc["warn"].append("Q%d 选项引用的题肢缺失：%s" % (n, oc["circled_missing"]))
            if not ks:
                sc["warn"].append("Q%d 未见答案键（answer_key_missing）" % n)
            elif len({k[0] for k in ks}) > 1:
                sc["hard_fail"].append("Q%d 答案键冲突：%s" % (n, ks))
        else:
            if not any(t in (LAYER_TITLE["E1"], LAYER_TITLE["E3"]) for t, _ in secs):
                sc["warn"].append("Q%d 未切到 E1/E3 块" % n)
        rrows = [r_ for r in results for r_ in r["residual"]
                 if ("Q%d" % n) in r_.get("qid_candidates", []) and (r is paper or r["layer"] not in ("paper", "teacher"))]
        chk = {"score": meta["score"], "choice": meta["choice"], "answer_keys": [k[0] for k in ks],
               "residual": sorted({r_["id"] for r_ in rrows}), "route": "__ROUTE__"}
        route_basis[n] = {"residual_kinds": sorted({r_["kind"] for r_ in rrows}),
                          "native_stem": any(t in (LAYER_TITLE["paper"], LAYER_TITLE["teacher"]) for t, _ in secs)}
        qid = "%s-Q%d" % (exam_id, n)
        md = _qmd(qid, exam_id, n, meta, secs, _guide(secs, paper, guide_note), chk)
        qout[qid] = md
    # 评分块落到不存在的题号
    for r in results:
        if r["layer"] in ("paper", "teacher"):
            continue
        orphan = [n for n in r["split"]["questions"] if n not in sq]
        if orphan:
            sc["warn"].append("%s 有题号块不在原卷题号内：%s" % (r["info"]["source_id"], orphan))
    sc["score_by_question"] = detail
    sc["score_total_resolved"] = total
    sc["score_unresolved"] = unresolved
    declared = [s for s in paper["split"]["sections"] if s.get("total")]
    sc["score_declared_total"] = sum(s["total"] for s in declared) if declared else None
    if not unresolved and nums and total != 100:
        sc["hard_fail"].append("分值合计 %d ≠ 100" % total)
    elif unresolved:
        if sc["score_declared_total"] == 100 and total < 100:
            sc["warn"].append("Q%s 分值未解析（已解析合计 %d），按卷首分部总分 100 核对" % (
                ",".join(map(str, unresolved)), total))
        else:
            sc["hard_fail"].append("Q%s 分值未解析，且卷首分部总分为 %s（不是 100），总分无法核对（已解析合计 %d）" % (
                ",".join(map(str, unresolved)), sc["score_declared_total"], total))
    sc["score_total_is_100"] = (not unresolved and total == 100)
    if dec_n and dec_n != len(nums):
        msg = "题数 %d 与卷首“本部分共N题”合计 %d 不符" % (len(nums), dec_n)
        if sc["score_total_is_100"] and hit:
            sc["warn"].append(msg + "；但题数与 %s 相符且逐题分值合计 100，疑原件卷首题数印误，须人工看原页" % "、".join(hit))
        else:
            sc["hard_fail"].append(msg)
    sc["choice_keys"] = sum(1 for n in nums if n in keys)
    sc["pages_unmapped"] = {r["info"]["source_id"]: r["selfcheck"]["pages_unmapped"] for r in results
                            if r["selfcheck"]["pages_unmapped"]}
    sc["residual_items"] = sum(len(r["residual"]) for r in results)
    sc["residual_by_kind"] = dict(collections.Counter(x["kind"] for r in results for x in r["residual"]))
    sc["page_types"] = dict(collections.Counter(p["page_type"] for r in results for p in r["pages"]))
    # 题级 MD 放卷目录；图片在 sources/<sid>/assets，链接改为相对卷目录；写盘前先查引用是否存在
    for qid in list(qout):
        qout[qid] = re.sub(r"\]\((assets/(S[0-9a-f]{12})/)", lambda m: "](sources/%s/%s" % (m.group(2), m.group(1)),
                           qout[qid])
    if write:
        miss = []
        for qid, md in qout.items():
            miss += asset_check(md, edir)
        sc["missing_asset_refs"] = sorted(set(miss))
        if miss:
            sc["hard_fail"].append("图片引用失效 %d 处" % len(set(miss)))
    # 分流建议（机器可读）：C=须模型对照原页（题面只有 OCR/扫描/乱码、切题存疑、或本题/整卷有硬失败）；
    # B=文字可直出但有图、表、框、文本框要看；A=纯文字层、无残留。自检全过只说明机械层没有发现问题。
    exam_level = [h for h in sc["hard_fail"] if not re.match(r"^(?:\S+：)?(?:\S+ )?Q\d+ ", h)]
    heavy = {"scan_page", "ocr_layer_page", "garbled_symbols", "split_uncertain", "missing_source", "source_unreadable"}
    routes = {}
    for n in nums:
        rb = route_basis.get(n, {})
        why = []
        if exam_level:
            why.append("整卷硬失败")
        if any(re.search(r"(?<!\d)Q%d(?!\d)" % n, h) for h in sc["hard_fail"]):
            why.append("本题硬失败")
        if not rb.get("native_stem"):
            why.append("题面无 native 文字层")
        hk = sorted(set(rb.get("residual_kinds", [])) & heavy)
        if hk:
            why.append("残留:" + "/".join(hk))
        if why:
            route = "C"
        elif rb.get("residual_kinds"):
            route, why = "B", ["残留:" + "/".join(rb["residual_kinds"])]
        else:
            route = "A"
        routes[n] = {"route": route, "why": why}
        qid = "%s-Q%d" % (exam_id, n)
        if qid in qout:
            qout[qid] = qout[qid].replace('"__ROUTE__"', json.dumps(route))
    sc["routes"] = routes
    sc["route_counts"] = dict(collections.Counter(v["route"] for v in routes.values()))
    if write:
        os.makedirs(edir, exist_ok=True)
        for qid, md in qout.items():
            with open(os.path.join(edir, qid + ".md"), "w", encoding="utf-8") as fh:
                fh.write(md)
        front = render_blocks(paper["split"]["front"], paper["rids"])
        with open(os.path.join(edir, "00_exam_front_matter.md"), "w", encoding="utf-8") as fh:
            fh.write("# %s 卷首（doc2md machine_extracted）\n\n%s\n" % (exam_id, re.sub(
                r"\]\((assets/(S[0-9a-f]{12})/)", lambda m: "](sources/%s/%s" % (m.group(2), m.group(1)), front)))
    return {"selfcheck": sc, "questions": qout}


# ---------------------------------------------------------------- 比对工具（供 bank_retro_audit 等复用）
WIDTH_MAP = {"．": ".", "，": ",", "：": ":", "；": ";", "（": "(", "）": ")", "？": "?", "！": "!", "“": '"', "”": '"',
             "‘": "'", "’": "'", "【": "[", "】": "]", "～": "~", "—": "-", "－": "-", "─": "-", "–": "-",
             "·": "·", "•": "·", "・": "·", "〈": "<", "〉": ">", "［": "[", "］": "]", "｛": "{", "｝": "}",
             "＋": "+", "＝": "=", "％": "%", "／": "/", "＜": "<", "＞": ">", "\u3000": " ", "＂": '"', "＇": "'",
             "﹒": ".", "。": "。", "、": "、", "∶": ":", "︰": ":",
             "«": "《", "»": "》"}  # 个别方正导出 PDF 把书名号映射成 « »，字形相同
for _i in range(10):
    WIDTH_MAP[chr(0xFF10 + _i)] = str(_i)
for _i in range(26):
    WIDTH_MAP[chr(0xFF21 + _i)] = chr(65 + _i)
    WIDTH_MAP[chr(0xFF41 + _i)] = chr(97 + _i)
WIDTH_TABLE = str.maketrans(WIDTH_MAP)
RE_STRIP_ANN = re.compile(r"〔[^〕]{0,80}〕|（原文如此[^）]{0,30}）|\(原文如此[^)]{0,30}\)|<!--.*?-->|!\[[^\]]*\]\([^)]*\)")
RE_TAGS = re.compile(r"</?(u|sup|sub|br|em|td|th|tr|table|thead|tbody|span|b|strong|mark|font|i|s|del|ins|p|div)(\s[^>]*)?/?>")


def norm_display(t):
    """去空白与标记（保留字形），作为可读的对齐底稿。"""
    t = RE_STRIP_ANN.sub("", t or "")
    t = RE_TAGS.sub("", t)
    t = t.replace("&nbsp;", "").replace("&lt;", "<").replace("&gt;", ">").replace("&amp;", "&")
    t = re.sub(r"\\([|*_`\[\]()#])", r"\1", t)
    t = re.sub(r"[*`#|>\s\u3000\u00a0\u200b\ufeff│┃┌┐└┘├┤┬┴┼─━═║╔╗╚╝]", "", t)
    return t


def canon(t):
    """全半角等价折叠（一字对一字，长度不变），用于匹配；真实不同的标点（，与、）仍保持不同。"""
    return fix_radicals(t.translate(WIDTH_TABLE))


class TextIndex:
    """把若干段文字（带来源标签）拼成一条可检索的规范化语料，支持 k-gram 定位与字符级对齐。"""

    def __init__(self, k=6):
        self.k = k
        self.disp, self.canon_s, self.segs = [], [], []  # segs: (start, end, label)
        self._len = 0
        self.grams = None

    def add(self, text, label):
        d = norm_display(text)
        if not d:
            return
        c = canon(d)
        self.segs.append((self._len, self._len + len(c), label))
        self.disp.append(d)
        self.canon_s.append(c)
        self._len += len(c)
        self.disp.append("\x00")
        self.canon_s.append("\x00")
        self._len += 1
        self.grams = None

    def build(self):
        self.D = "".join(self.disp)
        self.C = "".join(self.canon_s)
        g = collections.defaultdict(list)
        k = self.k
        C = self.C
        for i in range(len(C) - k + 1):
            g[C[i:i + k]].append(i)
        self.grams = g
        self._starts = [s[0] for s in self.segs]
        return self

    def label_at(self, pos):
        import bisect
        i = bisect.bisect_right(self._starts, pos) - 1
        return self.segs[max(0, i)][2] if self.segs else None

    def locate(self, n, allow=None):
        """返回最可能的起点（语料偏移）；allow(label)->bool 可限定来源。"""
        if self.grams is None:
            self.build()
        k = self.k if len(n) >= 12 else 4
        votes = collections.Counter()
        if k != self.k:
            i = self.C.find(n)
            while i >= 0 and len(votes) < 50:
                if allow is None or allow(self.label_at(i)):
                    votes[i] += len(n)
                i = self.C.find(n, i + 1)
            for j in range(0, len(n) - k + 1):
                sub = n[j:j + k]
                p = self.C.find(sub)
                cnt = 0
                while p >= 0 and cnt < 30:
                    if allow is None or allow(self.label_at(p)):
                        votes[p - j] += 1
                    p = self.C.find(sub, p + 1)
                    cnt += 1
        else:
            for j in range(0, len(n) - k + 1):
                for pos in self.grams.get(n[j:j + k], [])[:60]:
                    if allow is None or allow(self.label_at(pos)):
                        votes[pos - j] += 1
        if not votes:
            return None
        # 相邻偏移合并投票（容忍增删字造成的错位）
        best, bestv = None, -1
        for off, v in votes.items():
            s = v + votes.get(off - 1, 0) + votes.get(off + 1, 0) + votes.get(off - 2, 0) + votes.get(off + 2, 0)
            if s > bestv:
                best, bestv = off, s
        return best

    def align(self, n, allow=None, slack=12):
        """n 为 canon 后的行。返回 dict(matched, ratio, start, end, ops[(tag, a, b, a0, b0)])，未定位返回 None。"""
        import difflib
        if not n:
            return None
        i = self.C.find(n)
        while i >= 0 and allow is not None and not allow(self.label_at(i)):
            i = self.C.find(n, i + 1)
        if i >= 0:
            return {"matched": len(n), "ratio": 1.0, "start": i, "end": i + len(n), "ops": [], "longest": len(n)}
        off = self.locate(n, allow)
        if off is None:
            return None
        w0 = max(0, off - slack)
        w1 = min(len(self.C), off + len(n) + slack)
        win = self.C[w0:w1]
        cut = win.find("\x00", max(0, off - w0 + len(n) // 2))
        if cut > 0:
            win = win[:cut]
        cut0 = win.rfind("\x00", 0, max(0, off - w0 + len(n) // 2))
        if cut0 >= 0:
            win = win[cut0 + 1:]
            w0 += cut0 + 1
        sm = difflib.SequenceMatcher(None, n, win, autojunk=False)
        ops = sm.get_opcodes()
        eq = [o for o in ops if o[0] == "equal"]
        if not eq:
            return None
        matched = sum(o[2] - o[1] for o in eq)
        b_start = eq[0][3] - eq[0][1] if eq[0][1] > 0 else eq[0][3]
        out = []
        first_i = ops.index(eq[0])
        last_i = ops.index(eq[-1])
        for idx, (tag, a0, a1, b0, b1) in enumerate(ops):
            if tag == "equal":
                continue
            if idx < first_i or idx > last_i:
                if tag == "insert":
                    continue
                if tag == "replace":
                    tag, b0, b1 = "delete", b0, b0
            out.append((tag, a0, a1, w0 + b0, w0 + b1))
        return {"matched": matched, "ratio": matched / float(len(n)), "start": w0 + max(0, b_start),
                "end": w0 + eq[-1][4], "ops": out, "longest": max(o[2] - o[1] for o in eq)}


# ---------------------------------------------------------------- 写入护栏
def _profile_dirs():
    """所有书册 profile 的 paths.* 目录（文件路径取其所在目录，如 central_state）；冻结册另列。"""
    dirs, frozen = [], []
    try:
        sys.path.insert(0, HERE)
        from profile_lib import list_profiles, load_profile, resolve
        for name in list_profiles():
            prof = load_profile(name)
            for k, v in (prof.get("paths") or {}).items():
                if not isinstance(v, str) or not v or k == "root" or "{" in v:
                    continue
                d = str(resolve(prof, v))
                if os.path.splitext(d)[1]:          # 文件（.json/.md）→ 保护其所在目录
                    d = os.path.dirname(d)
                dirs.append(d)
                if prof.get("frozen") and k in ("workspace", "aux_workspace", "review_entry", "review_history"):
                    frozen.append((d, prof))
            for v in ((prof.get("sources") or {}).get("roots") or []):
                dirs.append(str(resolve(prof, v)))
    except Exception:
        pass
    return dirs, frozen


def _stat_key(p):
    try:
        st = os.stat(p)
        return (st.st_dev, st.st_ino)
    except OSError:
        return None


def _ancestors(p):
    out = []
    while True:
        out.append(p)
        q = os.path.dirname(p)
        if q == p:
            return out
        p = q


def path_inside(out, root):
    """out 是否落在 root 内（含相等）。三重判断：realpath 前缀、大小写折叠前缀、逐级祖先的 inode 比对
    （macOS APFS 默认不区分大小写，只改大小写的路径指向同一目录）。"""
    if not root:
        return False
    o = os.path.realpath(os.path.abspath(out))
    r = os.path.realpath(os.path.abspath(root))
    for a, b_ in ((o, r), (o.casefold(), r.casefold())):
        if a == b_ or a.startswith(b_.rstrip(os.sep) + os.sep):
            return True
    rk = _stat_key(r)
    if rk is None:
        return False
    for anc in _ancestors(o):
        k = _stat_key(anc)
        if k is not None and k == rk:
            return True
    return False


def _under_source_top(o):
    """桌面或原料根下任一“YYYY模拟题/历年高考题及细则/宝典制作原料”目录（以后的 2027模拟题 同样受保护）。"""
    for base in {DESKTOP, os.path.dirname(CANON_ROOT), RAW, os.path.join(CANON_ROOT, "00_共同资料/原材料")}:
        if not path_inside(o, base):
            continue
        ob = os.path.realpath(os.path.abspath(o))
        rb = os.path.realpath(base)
        rel = ob[len(rb):].lstrip(os.sep) if ob.casefold().startswith(rb.casefold()) else ""
        first = rel.split(os.sep)[0] if rel else ""
        if first and RE_SOURCE_TOP.match(first):
            return os.path.join(rb, first)
    return None


def guard_out_dir(out, inputs=()):
    """输出目录护栏。拒绝：题库、原料根与各“YYYY模拟题”目录、原件所在目录、各书册 profile 的 paths.*、
    Skill 目录、整个桌面（其中只有题库流水线 receipts/ 对 BANK_ACTOR=shared_writer 放行）。冻结册先走 frozen_guard。"""
    if not out:
        raise SystemExit("没有输出目录：请给 --out-dir 或设置环境变量 DOC2MD_OUT")
    o = os.path.realpath(os.path.abspath(out))
    dirs, frozen = _profile_dirs()
    for d, prof in frozen:
        if path_inside(o, d):
            from profile_lib import frozen_guard
            frozen_guard(prof, "写入 doc2md 输出")
    roots = {ROOT, CANON_ROOT}
    deny = []
    for r in roots:
        deny += [os.path.join(r, "DeepSeek_政治题库资料库_20260918"), os.path.join(r, "00_共同资料")]
    deny += [BANK, RAW, os.path.dirname(HERE), os.path.expanduser("~/.claude/skills")]
    deny += [os.path.dirname(os.path.abspath(p)) for p in inputs if p]
    for d in deny + dirs:
        if d and path_inside(o, d):
            raise SystemExit("拒绝写入：输出目录 %s 落在只读区 %s 内" % (o, d))
    top = _under_source_top(o)
    if top:
        raise SystemExit("拒绝写入：输出目录 %s 落在原料目录 %s 内" % (o, top))
    desk = {DESKTOP, os.path.dirname(CANON_ROOT), ROOT, CANON_ROOT}
    if any(path_inside(o, d) for d in desk):
        allowed = [os.path.join(r, "后勤管理/MD全库流水线_20260921/receipts") for r in roots]
        actor = os.environ.get("BANK_ACTOR", "")
        writer = ""
        try:
            writer = json.load(open(STATE, encoding="utf-8")).get("shared_writer", "")
        except Exception:
            pass
        if not (any(path_inside(o, a) for a in allowed) and actor and actor == writer):
            raise SystemExit("拒绝写入：输出目录 %s 在桌面内（项目、各册书稿、原料都在这里）。只有题库写权人"
                             "（BANK_ACTOR=state.json 的 shared_writer）可写 %s；其他人请用 scratchpad。"
                             % (o, allowed[0]))
    return o


# ---------------------------------------------------------------- 命令行
def ocr_exe():
    exe = shutil.which("ocr-vision") or os.path.expanduser("~/.local/bin/ocr-vision")
    return exe if os.path.isfile(exe) else None


def check_deps():
    """运行环境：PyMuPDF（fitz）与 python-pptx；缺哪个返回中文说明，齐全返回 None。"""
    miss = []
    for mod, pkg in (("fitz", "PyMuPDF"), ("pptx", "python-pptx"), ("lxml", "lxml")):
        try:
            __import__(mod)
        except Exception:
            miss.append(pkg)
    if miss:
        return ("当前解释器 %s（Python %s）缺少 %s。请改用 /usr/bin/python3（3.9，已装齐），不要装新包。"
                % (sys.executable, sys.version.split()[0], "、".join(miss)))
    return None


def main(argv=None):
    ap = argparse.ArgumentParser(description="原件→题级 Markdown 候选（确定性，只读原件，只写 --out-dir）。"
                                             "退出码：0 无硬失败；1 有硬失败（见 selfcheck.json）；2 参数/输入错误、拒绝写入或异常")
    ap.add_argument("file", nargs="?", help="单个原件路径（docx/pdf/pptx/doc/rtf）")
    ap.add_argument("--exam", help="按转换主清单转换一卷的全部来源（与 file 二选一）")
    ap.add_argument("--role", default="原卷", help="单件模式的来源角色：原卷/教师版/评分材料/参考答案/讲评…（默认原卷）")
    ap.add_argument("--exam-id", help="单件模式的卷号，用于题键命名（默认用 source_id）")
    ap.add_argument("--out-dir", default=DEFAULT_OUT,
                    help="输出根目录（必须给出；或设环境变量 DOC2MD_OUT。当前默认：%(default)s）")
    ap.add_argument("--ocr-dir", action="append", help="已有 OCR 文本目录（pNNN.ocr.txt）；默认查题库 evidence/<sid>/visual_transcript")
    ap.add_argument("--run-ocr", action="store_true", help="对缺 OCR 的扫描页用 ocr-vision 补跑（产物只写 out-dir；找不到 ocr-vision 时退出码 2）")
    ap.add_argument("--render-dpi", type=int, default=110, help="扫描页整页渲染 dpi（0=不渲染）")
    ap.add_argument("--quiet", action="store_true", help="只打印摘要一行")
    a = ap.parse_args(argv)
    if bool(a.file) == bool(a.exam):
        ap.print_usage()
        print("需要且只能给一个：原件路径 或 --exam")
        return 2
    dep = check_deps()
    if dep:
        print("运行环境不满足：%s" % dep)
        return 2
    if a.run_ocr and not ocr_exe():
        print("给了 --run-ocr，但找不到 ocr-vision（PATH 与 ~/.local/bin 都没有）；去掉 --run-ocr 或先装好再跑")
        return 2
    try:
        if a.file:
            if not os.path.isfile(a.file):
                print("找不到原件：%s" % a.file)
                return 2
            guard_out_dir(a.out_dir, [a.file])
            ocr_dirs = a.ocr_dir
            if ocr_dirs is None:
                ocr_dirs = [os.path.join(BANK, "evidence", sid_of(a.file), "visual_transcript")]
            res = convert_file(a.file, a.role, a.exam_id, a.out_dir, ocr_dirs=ocr_dirs, run_ocr_pages=a.run_ocr,
                               render_dpi=a.render_dpi)
            sc = res["selfcheck"]
            summ = {"输出": res["odir"], "页型": res["info"]["page_types"], "切出题块": len(res["split"]["questions"]),
                    "residual": len(res["residual"]), "硬失败": sc["hard_fail"], "警告数": len(sc["warn"]),
                    "耗时秒": res["info"]["seconds_total"]}
        else:
            task, srcs = exam_sources(a.exam)
            guard_out_dir(a.out_dir, [s["path"] for s in srcs if s["path"]])
            res = convert_exam(a.exam, a.out_dir, run_ocr_pages=a.run_ocr, render_dpi=a.render_dpi)
            sc = res["selfcheck"]
            summ = {"输出": os.path.join(a.out_dir, a.exam), "题面来源": sc.get("paper_role"), "题数": sc.get("questions"),
                    "应有题数": sc.get("expected_counts"),
                    "题号连续": sc.get("numbers_contiguous"), "分值合计": sc.get("score_total_resolved"),
                    "未解析分值": sc.get("score_unresolved"), "答案键": sc.get("choice_keys"),
                    "页型": sc.get("page_types"), "residual": sc.get("residual_items"),
                    "硬失败": sc["hard_fail"], "警告数": len(sc["warn"]), "耗时秒": sc.get("seconds")}
    except SystemExit as e:
        print(e)
        return 2
    except RuntimeError as e:
        print("转换失败：%s" % e)
        return 2
    except Exception as e:  # 未预料的异常也退 2，退出码 1 只留给“已写 selfcheck.json 且有硬失败”
        print("转换异常（%s: %s）；未完成的产物不可用。" % (type(e).__name__, e))
        return 2
    print(json.dumps(summ, ensure_ascii=False))
    if not a.quiet:
        print("摘要：%s；硬失败 %d 项，警告 %d 项；产物证据等级 machine_extracted。自检全过只说明机械层没有发现问题，"
              "不等于内容已核验。" % ("单件转换完成" if a.file else "整卷转换完成", len(sc["hard_fail"]), len(sc["warn"])))
    return 1 if sc["hard_fail"] else 0


if __name__ == "__main__":
    sys.exit(main())
