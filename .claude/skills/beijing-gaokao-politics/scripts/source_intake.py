#!/usr/bin/env python3
"""Fingerprint supplied source packages; cache text without changing original files.

This is discovery/extraction, not evidence acceptance, OCR or human verification.
"""
import argparse, hashlib, json, zipfile
from pathlib import Path
from lxml import etree as E
EXTS={'.pdf','.pptx','.docx','.xlsx','.xls','.doc','.ppt','.zip','.png','.jpg','.jpeg','.tif','.tiff'}

def digest(p):
 h=hashlib.sha256()
 with p.open('rb') as f:
  for b in iter(lambda:f.read(1024*1024),b''):h.update(b)
 return h.hexdigest()

def extract(p):
 ext=p.suffix.lower();units=[];extra={}
 if ext=='.pdf':
  from pypdf import PdfReader
  reader=PdfReader(p)
  units=[{'unit':i+1,'text':page.extract_text() or ''} for i,page in enumerate(reader.pages)]
 elif ext=='.pptx':
  from pptx import Presentation
  pres=Presentation(p)
  for i,s in enumerate(pres.slides):
   ts=[]
   def visit(shapes):
    for shape in shapes:
     if shape.has_text_frame:ts.append(shape.text_frame.text)
     if shape.has_table:
      ts.extend(' | '.join(c.text for c in row.cells) for row in shape.table.rows)
     if hasattr(shape,'shapes'):visit(shape.shapes)
   visit(s.shapes)
   notes=s.notes_slide.notes_text_frame.text if s.has_notes_slide else ''
   units.append({'unit':i+1,'text':'\n'.join(ts),'notes':notes})
  with zipfile.ZipFile(p) as z:extra={'media_members':[n for n in z.namelist() if n.startswith('ppt/media/')],'embedded_members':[n for n in z.namelist() if n.startswith('ppt/embeddings/')]}
 elif ext=='.docx':
  with zipfile.ZipFile(p) as z:
   r=E.fromstring(z.read('word/document.xml'));ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
   units=[{'unit':i+1,'text':''.join(n.xpath('.//w:t/text()',namespaces=ns))} for i,n in enumerate(r.findall('w:body/*',ns))]
   extra={'media_members':[n for n in z.namelist() if n.startswith('word/media/')]}
 elif ext=='.zip':
  with zipfile.ZipFile(p) as z:extra={'archive_members':[{'name':i.filename,'bytes':i.file_size} for i in z.infolist()]}
 return {'extractor_revision':1,'units':units,'unit_count':len(units),'empty_text_units':[u['unit'] for u in units if not u['text'].strip()],**extra,'human_verified':False,'note':'Text extraction can omit images, tables, hidden content and encoding errors. Empty text is not proof that a page or question is absent.'}

def scan(inputs,cache,previous=None):
 cache.mkdir(parents=True,exist_ok=True)
 prev=json.loads(Path(previous).read_text()) if previous else {'files':[]}
 old={x['path']:x for x in prev['files']};oldhash={x['sha256'] for x in prev['files']};files=[];seen=set()
 for supplied in inputs:
  supplied=Path(supplied).resolve()
  if not supplied.exists():raise FileNotFoundError(supplied)
  for p in sorted(supplied.rglob('*') if supplied.is_dir() else [supplied]):
   if not p.is_file() or p.suffix.lower() not in EXTS or p.name.startswith('~$'):continue
   p=p.resolve()
   if p in seen:continue
   if p==cache.resolve() or cache.resolve() in p.parents:continue
   seen.add(p);sha=digest(p);key=str(p);cached=cache/(sha+'.json')
   status='unchanged' if key in old and old[key]['sha256']==sha else 'same_bytes_other_path' if sha in oldhash else 'changed' if key in old else 'new'
   hit=cached.exists()
   if not hit:
    try:data=extract(p)
    except Exception as e:data={'extractor_revision':1,'units':[],'human_verified':False,'extraction_error':str(e)}
    data.update({'source_sha256':sha,'source_suffix':p.suffix.lower()});cached.write_text(json.dumps(data,ensure_ascii=False,indent=2))
   files.append({'path':key,'bytes':p.stat().st_size,'sha256':sha,'status':status,'cache':str(cached),'cache_reused':hit,'source_identity_and_evidence_status':'not_inferred_from_filename'})
 return {'schema':1,'inputs':[str(Path(p).resolve()) for p in inputs],'files':files,'counts':{s:sum(x['status']==s for x in files) for s in ['new','changed','unchanged','same_bytes_other_path']},'purpose':'Source intake only; root must verify full exam identity and question-level formal scoring before insertion.'}

def main():
 p=argparse.ArgumentParser(description=__doc__);p.add_argument('input',nargs='+',type=Path);p.add_argument('--cache',type=Path,required=True);p.add_argument('--previous',type=Path);p.add_argument('--report',type=Path,required=True);a=p.parse_args()
 assert not a.report.exists(),'Choose a new report; preserve previous intake record'
 a.report.parent.mkdir(parents=True,exist_ok=True);r=scan(a.input,a.cache,a.previous);a.report.write_text(json.dumps(r,ensure_ascii=False,indent=2));print(json.dumps(r['counts'],ensure_ascii=False))
if __name__=='__main__':main()
