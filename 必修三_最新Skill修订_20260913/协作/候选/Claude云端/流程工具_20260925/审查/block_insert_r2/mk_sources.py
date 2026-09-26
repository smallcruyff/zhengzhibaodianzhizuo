import sys
sys.dont_write_bytecode = True
import zipfile, io, re
from pathlib import Path
import docx
from common import WORK

def make_base(path, cell_texts):
    d = docx.Document()
    d.add_paragraph('原件前一段')
    t = d.add_table(rows=1, cols=len(cell_texts))
    for i, s in enumerate(cell_texts):
        t.rows[0].cells[i].text = s
    d.add_paragraph('原件后一段')
    d.save(path)

def patch(path, doc_fn, rels_fn=None, extra_members=None):
    src = Path(path)
    buf = io.BytesIO()
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == 'word/document.xml':
                data = doc_fn(data.decode('utf-8')).encode('utf-8')
            elif info.filename == 'word/_rels/document.xml.rels' and rels_fn:
                data = rels_fn(data.decode('utf-8')).encode('utf-8')
            zout.writestr(info, data)
        for n, b in (extra_members or {}).items():
            zout.writestr(n, b)
    src.write_bytes(buf.getvalue())

# A: 表格内只有一个外部超链接，无图
a = WORK / 'src_hyperlink_only.docx'
make_base(a, ['甲单元格', '乙单元格'])
def docA(x):
    return x.replace('<w:t>甲单元格</w:t></w:r>',
        '<w:t>甲单元格</w:t></w:r><w:hyperlink r:id="rIdHL1"><w:r><w:t>链接文字</w:t></w:r></w:hyperlink>', 1)
def relsA(x):
    return x.replace('</Relationships>',
        '<Relationship Id="rIdHL1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="http://example.com/a" TargetMode="External"/></Relationships>')
patch(a, docA, relsA)

# B: 表格内一张 VML 老式图片（w:pict/v:shape/v:imagedata r:id）
import fitz
pm = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 40, 30), 0)
pm.clear_with(200)
png = pm.tobytes('png')
b = WORK / 'src_vml_in_table.docx'
make_base(b, ['丙单元格', '丁单元格'])
def docB(x):
    if 'xmlns:v=' not in x:
        x = x.replace('<w:document ', '<w:document xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office" ', 1)
    return x.replace('<w:t>丙单元格</w:t></w:r>',
        '<w:t>丙单元格</w:t></w:r><w:r><w:pict><v:shape id="_x0000_i1025" type="#_x0000_t75" style="width:30pt;height:22.5pt"><v:imagedata r:id="rIdVML1" o:title=""/></v:shape></w:pict></w:r>', 1)
def relsB(x):
    return x.replace('</Relationships>',
        '<Relationship Id="rIdVML1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/vml1.png"/></Relationships>')
patch(b, docB, relsB, {'word/media/vml1.png': png})
print(a, b)
