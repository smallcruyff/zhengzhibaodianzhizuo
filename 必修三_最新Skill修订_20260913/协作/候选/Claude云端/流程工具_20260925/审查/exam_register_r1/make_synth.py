import sys
sys.dont_write_bytecode = True
from pathlib import Path
import docx
S = Path('/tmp/claude-0/-home-user-zhengzhibaodianzhizuo/2415817c-761f-5954-bb79-9bc9fa3464ca/scratchpad/review_exam_register_r1/synth')
S.mkdir(parents=True, exist_ok=True)
QS = [
 '1．某市推行“接诉即办”改革，市民热线诉求当天响应、限期办结，并将办理结果纳入部门考核。这一做法①坚持以人民为中心的发展思想②说明政府职能已经由管理转向服务③有利于提高政府的公信力和执行力④表明群众可以直接决定政府的施政方向A．①②B．①③C．②④D．③④',
 '2．某县探索“党建引领、村民议事、乡贤助力”的乡村治理模式，村内公共事务由村民议事会讨论决定。这一模式①丰富了基层民主的实现形式②改变了村民委员会的性质③有利于调动村民参与治理的积极性④说明村民议事会是基层政权组织A．①③B．①④C．②③D．②④',
 '3．某省人大常委会在制定地方性法规过程中，通过基层立法联系点广泛收集群众意见，并对意见采纳情况予以反馈。这体现了①全过程人民民主的实践②地方人大享有国家立法权③科学立法、民主立法的要求④人民群众直接行使立法权A．①③B．①④C．②③D．②④',
 '4．某区检察院针对校园周边食品安全问题发出检察建议，督促相关行政机关依法履职并开展专项整治。这说明①检察机关是国家的法律监督机关②检察机关可以直接行使行政处罚权③检察建议有助于推动依法行政④检察机关领导行政机关开展工作A．①③B．①④C．②③D．②④',
 '5．某市政协围绕“老旧小区改造”开展专题协商，委员们深入调研并提出多项建议，被有关部门吸收采纳。这表明①人民政协是专门协商机构②人民政协是国家权力机关③政协协商有助于科学民主决策④政协委员可以代替政府作出决策A．①③B．①④C．②③D．②④',
 '6．某地法院设立“共享法庭”，把诉讼服务延伸到村社网格，就地化解矛盾纠纷。这一做法①便利群众参与诉讼②改变了法院的审判职能③推动矛盾纠纷源头化解④意味着调解可以替代审判A．①③B．①④C．②③D．②④',
]
def mk(name, header_lines, qs=QS, q1_override=None):
    d = docx.Document()
    for h in header_lines:
        d.add_paragraph(h)
    for i, q in enumerate(qs):
        d.add_paragraph(q1_override if (i == 0 and q1_override) else q)
    d.save(str(S / name))
# S1 no date at all
mk('s1_nodate.docx', ['北京市某区高三思想政治期中练习', '思想政治'])
# S2 高三 header + 高二 in Q1 body
mk('s2_g3_header_g2_body.docx', ['海淀区2024—2025学年第一学期期中练习', '高三思想政治', '2024.11'],
   q1_override='1．某校高二学生开展“模拟政协”活动，围绕校园周边交通拥堵问题撰写提案并提交区政协。这一活动①有利于增强学生的公民意识②说明学生可以直接参与国家立法③体现了人民政协履行职能的开放性④表明政协提案具有法律效力A．①③B．①④C．②③D．②④')
# S3 genuine 高二 paper
mk('s3_grade2.docx', ['海淀区2024—2025学年第一学期期中练习', '高二思想政治', '2024.11'])
# S4 website-style title (exam calendar year, no month) + question-body date from previous year
mk('s4_webtitle_bodydate.docx', ['2025北京朝阳高三（上）期中', '政  治', '（考试时间90分钟 满分100分）'],
   q1_override='1．2024年10月，某市出台《促进民营经济高质量发展若干措施》，从市场准入、要素获取、公平执法等方面提出具体举措。这表明①我国坚持“两个毫不动摇”②民营经济是国民经济的主导力量③优化营商环境有利于激发市场活力④政府直接配置民营企业的生产要素A．①③B．①④C．②③D．②④')
# S5 header with full-width digits and PUA dot like real CY/SY papers
mk('s5_fullwidth.docx', ['北京市朝阳区高三年级第一学期质量检测', '思想政治 ２０２４\U00100170 １１', '（考试时间９０分钟 满分１００分）'])
print('ok')
