#!/usr/bin/env python3
"""exam_register r3 对抗探针（只读真实输入；合成件与报告只写 scratchpad；结果记到本目录 probe_cli_r3.json）。"""
import json, os, shutil, subprocess, sys, time
sys.dont_write_bytecode = True
from pathlib import Path
C = Path(__file__).resolve().parents[2]
TOOL = C / 'exam_register.py'
PROF = C / 'profiles_cloud' / 'bixiu3.json'
S = Path(os.environ.get('R3_SCRATCH', '/tmp/claude-0/-home-user-zhengzhibaodianzhizuo/2415817c-761f-5954-bb79-9bc9fa3464ca/scratchpad/review_exam_register_r3'))
SY = S / 'synth'
REPO = C.parents[4]
RAW = REPO / '00_共同资料' / '原材料'
out = {}

def run(name, args, extra_env=None, pyflags=('-B',)):
    env = dict(os.environ)
    env.update(extra_env or {})
    t0 = time.time()
    p = subprocess.run([sys.executable, *pyflags, str(TOOL)] + args, capture_output=True, text=True, env=env, timeout=400)
    try:
        d = json.loads(p.stdout)
    except Exception:
        d = None
    rec = {'args': args, 'exit': p.returncode, 'seconds': round(time.time() - t0, 1), 'stderr': p.stderr.strip()[-600:]}
    if d is not None:
        rec['summary'] = {k: d.get(k) for k in ('cohort_year', 'cohort_status', 'grade2', 'suggestion', 'chain_kind',
                                                  'rename_plan', 'redirects', 'reused_ids', 'id_mismatch')}
        rec['duplicates'] = [(a['exam_id'], a['kind'], a['matched'], a['total']) for a in d.get('duplicates') or []]
        rec['issues'] = d.get('issues')
    else:
        rec['stdout_head'] = p.stdout[:300]
    out[name] = rec
    print('==', name, 'exit', p.returncode, json.dumps({k: rec.get(k) for k in ('summary', 'duplicates')}, ensure_ascii=False)[:900])
    if rec.get('issues'):
        for i in rec['issues']:
            print('    ISSUE', i[:220])
    if p.stderr.strip():
        print('    STDERR', p.stderr.strip()[-300:])
    return rec

def make_docx(name, header, qs):
    import docx
    d = docx.Document()
    for h in header:
        d.add_paragraph(h)
    for i, q in enumerate(qs):
        d.add_paragraph(f'{i + 1}．{q}')
    p = SY / name
    d.save(str(p))
    return p

FILL = [
 '某市推行“接诉即办”改革，市民热线诉求当天响应、限期办结，并将办理结果纳入部门考核。这一做法①坚持以人民为中心②说明政府职能已经由管理转向服务③有利于提高政府的公信力④表明群众可以直接决定政府的施政方向A．①②B．①③C．②④D．③④',
 '某县探索“党建引领、村民议事、乡贤助力”的乡村治理模式，村内公共事务由村民议事会讨论决定。这一模式①丰富了基层民主的实现形式②改变了村民委员会的性质③有利于调动村民参与治理的积极性④说明村民议事会是基层政权组织A．①③B．①④C．②③D．②④',
 '某省人大常委会在制定地方性法规过程中，通过基层立法联系点广泛收集群众意见，并对意见采纳情况予以反馈。这体现了①全过程人民民主的实践②地方人大享有国家立法权③科学立法民主立法的要求④人民群众直接行使立法权A．①③B．①④C．②③D．②④',
 '某区检察院针对校园周边食品安全问题发出检察建议，督促相关行政机关依法履职并开展专项整治。这说明①检察机关是国家的法律监督机关②检察机关可以直接行使行政处罚权③检察建议有助于推动依法行政④检察机关领导行政机关开展工作A．①③B．①④C．②③D．②④',
]

def main():
    if SY.exists():
        shutil.rmtree(SY)
    SY.mkdir(parents=True)
    P = ['--profile', str(PROF)]
    hd24 = RAW / '2024模拟题/2024海淀期中/试卷/试卷.pdf'
    hd25 = RAW / '2025模拟题/2025各区期末/2025海淀期中/试卷/试卷.pdf'
    cy24a = RAW / '2024模拟题/2024朝阳期中/试卷/试卷.pdf'
    cy24b = RAW / '2024模拟题/2024朝阳期中/试卷/补充材料/202411朝阳高三政治 期中1试题(1).pdf'
    # N1 乱码文字层：2024.11 海淀期中原件（即 ledger 里与 BJ-2025-HD-QIZHONG 同卷的那份）
    run('N1_garbled_textlayer_hd2024qizhong', ['check', '--files', str(hd24), '--region', '海淀', '--stage', '期中'] + P)
    # N2 多文件分母稀释：同一原卷的两份同 SHA 副本（exam_files.csv 两行都标“原卷”）
    run('N2a_cy24_single', ['check', '--files', str(cy24a), '--region', '朝阳', '--stage', '期中'] + P)
    run('N2b_cy24_two_identical_copies', ['check', '--files', str(cy24a), str(cy24b), '--region', '朝阳', '--stage', '期中'] + P)
    run('N2c_hd25_same_file_twice', ['check', '--files', str(hd25), str(hd25), '--region', '海淀', '--stage', '期中'] + P)
    # N3 --files 新卷改号链：reused_ids 漏掉本卷将接手的卷号
    new_hd23 = make_docx('new_hd_2022_11_qizhong.docx', ['海淀区2022—2023学年第一学期期中练习', '高三思想政治', '2022.11'], FILL)
    run('N3a_files_chain_hd2023', ['check', '--files', str(new_hd23), '--region', '海淀', '--stage', '期中'] + P)
    new_cy24 = make_docx('new_cy_2023_11_qizhong.docx', ['北京市朝阳区2023～2024学年度第一学期期中质量检测', '高三思想政治试卷', '2023.11'], FILL)
    run('N3b_files_chain_cy2024', ['check', '--files', str(new_cy24), '--region', '朝阳', '--stage', '期中'] + P)
    # N4 --debug 放在子命令前（文档称“可放在子命令前后”）
    run('N4a_debug_before_subcmd', ['--debug', 'check', '--exam-id', 'BJ-2024-HD-QIZHONG', '--raw-root', '/no/such'] + P)
    run('N4b_debug_after_subcmd', ['check', '--debug', '--exam-id', 'BJ-2024-HD-QIZHONG', '--raw-root', '/no/such'] + P)
    # N5 守卫：报告落候选 审查/（不在 构建/ 下）
    r = C / '审查' / 'exam_register_r3' / 'should_not_exist.json'
    rec = run('N5a_report_into_C_审查', ['check', '--exam-id', 'BJ-2026-HD-QIZHONG', '--report', str(r)] + P)
    rec['wrote'] = r.exists()
    if r.exists():
        r.unlink()
    # N5b 守卫：scratchpad 里的符号链接目录指向题库 indexes/
    link = S / 'link_to_bank_indexes'
    if link.is_symlink() or link.exists():
        link.unlink()
    os.symlink(str(REPO / 'DeepSeek_政治题库资料库_20260918' / 'indexes'), str(link))
    tgt = REPO / 'DeepSeek_政治题库资料库_20260918' / 'indexes' / 'r3_probe.json'
    rec = run('N5b_report_via_symlink_into_bank', ['check', '--exam-id', 'BJ-2026-HD-QIZHONG', '--report', str(link / 'r3_probe.json')])
    rec['wrote'] = tgt.exists()
    if tgt.exists():
        tgt.unlink()
    link.unlink()
    # N5c 守卫：非 .json 后缀
    run('N5c_report_bad_suffix', ['check', '--exam-id', 'BJ-2026-HD-QIZHONG', '--report', str(S / 'x.txt')] + P)
    # N6 不加 -B 运行一次，看 SK/候选目录是否留 __pycache__
    run('N6_no_B_flag', ['check', '--exam-id', 'BJ-2026-HD-QIZHONG'] + P, pyflags=())
    sk = REPO / '.claude/skills/beijing-gaokao-politics/scripts'
    out['N6_no_B_flag']['sk_pycache'] = [str(p) for p in sk.rglob('__pycache__')]
    out['N6_no_B_flag']['c_pycache'] = [str(p) for p in C.glob('__pycache__/exam_register*')]
    print('   SK __pycache__:', out['N6_no_B_flag']['sk_pycache'], 'C exam_register pyc:', out['N6_no_B_flag']['c_pycache'])
    (Path(__file__).parent / 'probe_cli_r3.json').write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding='utf-8')

if __name__ == '__main__':
    main()
