#!/usr/bin/env python3
"""r3 探针：新卷 20 题里只复用某张高考卷 2 题，但第 4 题题号排版异常（“4 ”无句点），--files 切题只切出
3 块（第 3 块吞掉 Q3–Q20）；看工具是否把“少数题复用”判成 duplicate_registration（fraction_reliable=True）。
合成件只写 scratchpad；只读题库。"""
import sys, os, json, subprocess
sys.dont_write_bytecode = True
from pathlib import Path
C = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(C))
import exam_register as er  # noqa
S = Path('/tmp/claude-0/-home-user-zhengzhibaodianzhizuo/2415817c-761f-5954-bb79-9bc9fa3464ca/scratchpad/review_exam_register_r3/synth')
S.mkdir(parents=True, exist_ok=True)
bank = er.REPO_ROOT / 'DeepSeek_政治题库资料库_20260918'
T, _ = er.load_bank_alias_tables(bank)
def stem(eid, n):
    p = bank / 'questions' / eid / f'{eid}-Q{n}.md'
    s, b = er.extract_stem_from_md(p.read_text(encoding='utf-8'), T)
    return s.strip().replace('\n', '')
r1, r2 = stem('BJ-2025-BJ-GAOKAO', 1), stem('BJ-2025-BJ-GAOKAO', 5)
FILL = ['某市推行“接诉即办”改革，市民热线诉求当天响应、限期办结，并将办理结果纳入部门考核。这一做法①坚持以人民为中心②说明政府职能已经由管理转向服务③有利于提高政府的公信力④表明群众可以直接决定政府的施政方向A．①②B．①③C．②④D．③④',
        '某县探索“党建引领、村民议事、乡贤助力”的乡村治理模式，村内公共事务由村民议事会讨论决定。这一模式①丰富了基层民主的实现形式②改变了村民委员会的性质③有利于调动村民参与治理的积极性④说明村民议事会是基层政权组织A．①③B．①④C．②③D．②④',
        '某省人大常委会在制定地方性法规过程中，通过基层立法联系点广泛收集群众意见，并对意见采纳情况予以反馈。这体现了①全过程人民民主的实践②地方人大享有国家立法权③科学立法民主立法的要求④人民群众直接行使立法权A．①③B．①④C．②③D．②④']
import docx
d = docx.Document()
for h in ['海淀区2026—2027学年第一学期期末练习', '高三思想政治', '2027.01']:
    d.add_paragraph(h)
qs = [r1] + [FILL[i % 3] + f'（第{i+2}题）' for i in range(18)] + [r2]
for i, q in enumerate(qs):
    n = i + 1
    d.add_paragraph((f'{n} ' if n == 4 else f'{n}．') + q)   # 第 4 题题号漏了句点
out = S / 'poor_split_2_reused_of_20.docx'
d.save(str(out))
p = subprocess.run([sys.executable, '-B', str(C / 'exam_register.py'), 'check', '--files', str(out), '--region', '海淀',
                    '--stage', '期末', '--profile', str(C / 'profiles_cloud' / 'bixiu3.json')], capture_output=True, text=True)
res = json.loads(p.stdout)
print('exit', p.returncode, 'questions_in_paper', len(qs))
print(json.dumps({'duplicates': res['duplicates'], 'issues': res['issues']}, ensure_ascii=False, indent=1))
