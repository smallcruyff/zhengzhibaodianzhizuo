#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""阶段0 能力验证：ocr-vision / soffice doc->docx / pdftoppm 渲染。

规范四.1："先测试能力和实际输出，不能假定某个库天然支持所有图片、备注和复杂表格。"
结果写入 validation/00_能力验证.json，供后续引用。
"""
import json
import os
import shutil
import sys
import time

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import _paths as P

OUT = {}
V = P.VAL_DIR
T = os.path.join(P.TMP_DIR, "capcheck")
P.ensure_dirs(V, T)


def log(k, v):
    OUT[k] = v
    print("[%s] %s" % (k, json.dumps(v, ensure_ascii=False)[:900]))


# ---------------------------------------------------------------- 样本
SCAN_PDF = os.path.join(P.PRIMARY_ROOT, "2026模拟题/2026各区一模/2026东城一模/试卷/补充材料/2026东城一模 原卷扫描版.pdf")
DOC_SAMPLE = os.path.join(P.PRIMARY_ROOT, "2026模拟题/2026各区二模/2026石景山二模/细则/石景山区高三政治第二次模拟考试答案评分细则(1).doc")

log("sample_scan_pdf", {"path": SCAN_PDF, "exists": os.path.isfile(SCAN_PDF)})
log("sample_doc", {"path": DOC_SAMPLE, "exists": os.path.isfile(DOC_SAMPLE)})

# ---------------------------------------------------------------- 1. PDF 探针
rc, out, err = P.run([P.PDFINFO, SCAN_PDF], timeout=60)
info = {}
for line in out.splitlines():
    if ":" in line:
        k, _, v = line.partition(":")
        info[k.strip()] = v.strip()
pages = int(info.get("Pages", "0") or 0)
log("pdfinfo", {"rc": rc, "Pages": pages, "Page size": info.get("Page size"),
                "Encrypted": info.get("Encrypted"), "Producer": info.get("Producer")})

rc, txt, err = P.run([P.PDFTOTEXT, "-layout", SCAN_PDF, "-"], timeout=120)
nchar = len(txt.strip())
log("pdftotext_scan", {"rc": rc, "chars": nchar,
                       "chars_per_page": round(nchar / pages, 1) if pages else None,
                       "verdict": "扫描/纯图" if pages and nchar / pages < 100 else "有文字层"})

# ---------------------------------------------------------------- 2. pdftoppm 渲染计时
r1 = os.path.join(T, "scan")
shutil.rmtree(r1, ignore_errors=True)
os.makedirs(r1)
t0 = time.time()
rc, out, err = P.run([P.PDFTOPPM, "-r", "300", "-png", "-f", "1", "-l", "1", SCAN_PDF, os.path.join(r1, "p")], timeout=180)
dt = time.time() - t0
made = sorted(os.listdir(r1))
sz = sum(os.path.getsize(os.path.join(r1, f)) for f in made)
log("pdftoppm_300dpi", {"rc": rc, "seconds": round(dt, 2), "files": made,
                        "bytes": sz, "kb": round(sz / 1024, 1),
                        "est_full_pages_sec": round(dt * pages, 1) if pages else None})
PNG1 = os.path.join(r1, made[0]) if made else None

# ---------------------------------------------------------------- 3. ocr-vision 端到端
if PNG1:
    r2 = os.path.join(T, "ocr")
    shutil.rmtree(r2, ignore_errors=True)
    os.makedirs(r2)
    t0 = time.time()
    rc, out, err = P.run([P.OCR_VISION, "--out", r2, "--langs", "zh-Hans,en-US", "--scale", "3.0", PNG1], timeout=300)
    dt = time.time() - t0
    produced = sorted(os.listdir(r2))
    text = ""
    for f in produced:
        fp = os.path.join(r2, f)
        if os.path.isfile(fp):
            with open(fp, "r", encoding="utf-8", errors="replace") as fh:
                text += fh.read()
    log("ocr_vision", {"rc": rc, "seconds": round(dt, 2), "produced": produced,
                       "chars": len(text), "stderr": err[:300],
                       "head": text[:600]})

# ---------------------------------------------------------------- 4. soffice doc -> docx
r3 = os.path.join(T, "soffice")
shutil.rmtree(r3, ignore_errors=True)
os.makedirs(r3)
env = P.env_with_fontconfig()
t0 = time.time()
rc, out, err = P.run([P.SOFFICE, "--headless", "--norestore", "--convert-to", "docx",
                      "--outdir", r3, DOC_SAMPLE], timeout=300, env=env)
dt = time.time() - t0
made = sorted(os.listdir(r3))
log("soffice_doc2docx", {"rc": rc, "seconds": round(dt, 2), "produced": made,
                         "stdout": out[:300], "stderr": err[:300],
                         "fontconfig": P.fontconfig_file()})

DOCX_OUT = os.path.join(r3, made[0]) if made else None
if DOCX_OUT and DOCX_OUT.endswith(".docx"):
    try:
        from docx import Document
        d = Document(DOCX_OUT)
        paras = [p.text for p in d.paragraphs if p.text.strip()]
        log("docx_after_convert", {"paragraphs": len(d.paragraphs), "tables": len(d.tables),
                                   "nonempty_paras": len(paras), "head": paras[:8]})
    except Exception as e:
        log("docx_after_convert", {"error": repr(e)})

with open(os.path.join(V, "00_能力验证.json"), "w", encoding="utf-8") as fh:
    json.dump(OUT, fh, ensure_ascii=False, indent=2)
print("\nWROTE", os.path.join(V, "00_能力验证.json"))
